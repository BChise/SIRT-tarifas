"""
Lee datos/pruebas.xlsx (una hoja por servicio) y calcula el estado de cada fila a partir
de los CSV de resultados: Pendiente → Registrada → Definitivo (o Error).
"""

import csv
import re
from pathlib import Path

from app import catalogo

ARCHIVO = catalogo.PROYECTO / "datos" / "pruebas.xlsx"
RESULTADOS = catalogo.RESULTADOS

COLUMNAS = {
    "n": "N°", "nombre": "Nombre de la tarifa", "codigo": "Código de servicio", "alcance": "Alcance",
    "lote": "Lote", "definitivo": "Registro definitivo", "accion": "Acción",
    "parametros": "Parámetros de la acción", "campos": "Campos a modificar", "observaciones": "Observaciones",
}


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def leer_filas() -> list[dict]:
    """Filas del bloque «TARIFAS ORIGINALES» de cada hoja de servicio (la hoja se llama
    como el código: TF, IF, SMP…). Las columnas después de «Parámetros de la acción» son
    los atributos de la tarifa."""
    from openpyxl import load_workbook
    if not ARCHIVO.exists():
        raise FileNotFoundError(f"No existe {ARCHIVO}")
    libro = load_workbook(ARCHIVO, data_only=True)
    servicios = {s["codigo"]: s for s in catalogo.servicios()}
    filas = []
    for hoja in libro.worksheets:
        codigo = hoja.title.strip().upper()
        if codigo not in servicios:
            continue
        servicio = servicios[codigo]
        filas_hoja = list(hoja.iter_rows(values_only=True))
        inicio = next((i for i, f in enumerate(filas_hoja)
                       if f and _texto(f[0]) == "N°" and len(f) > 1 and _texto(f[1]) == COLUMNAS["nombre"]), None)
        if inicio is None:
            continue
        cabecera = [_texto(c) for c in filas_hoja[inicio]]
        indice = {clave: cabecera.index(titulo) for clave, titulo in COLUMNAS.items() if titulo in cabecera}
        i_alcance = cabecera.index("Alcance") if "Alcance" in cabecera else None
        primer_atributo = (indice.get("parametros", 5)) + 1
        for numero, fila in enumerate(filas_hoja[inicio + 1:], start=inicio + 2):
            if not fila or not any(v not in (None, "") for v in fila):
                break                                    # fin del bloque 1
            datos = {clave: _texto(fila[i]) if i < len(fila) else "" for clave, i in indice.items()}
            if not datos.get("nombre"):
                continue
            datos.update(codigo=codigo, hoja=hoja.title, fila_excel=numero, servicio=servicio["nombre"],
                         archivo=servicio["archivo"], lote=datos.get("lote") or "DEMO",
                         alcance=_texto(fila[i_alcance]) if i_alcance is not None else "",
                         campos=datos.get("campos", ""))
            datos["atributos"] = {cabecera[j]: _texto(fila[j]) for j in range(primer_atributo, len(cabecera))
                                  if j < len(fila) and cabecera[j]}
            datos["quiere_definitivo"] = datos.get("definitivo", "").lower().startswith("s")
            datos["problemas"] = []
            if not re.fullmatch(r"[A-Za-z0-9_-]{1,30}", datos["lote"]):
                datos["problemas"].append("lote inválido")
            filas.append(datos)
    nombres = [f["nombre"] for f in filas]
    for f in filas:
        if nombres.count(f["nombre"]) > 1:
            f["problemas"].append("nombre repetido en el Excel")
    return filas


def _leer(ruta: Path) -> list[dict]:
    try:
        with ruta.open("r", newline="", encoding="utf-8-sig") as entrada:
            return list(csv.DictReader(entrada, delimiter=";"))
    except Exception:
        return []


def con_estado(filas: list[dict]) -> list[dict]:
    """Agrega a cada fila: estado (Pendiente / Registrada / Definitivo / Error), detalle y fecha."""
    for f in filas:
        f.update(estado="Pendiente", detalle="", fecha="", id_registro="")
        if not f["codigo"]:
            continue
        registro = [r for r in _leer(RESULTADOS / f["codigo"] / f"resultados_{f['lote']}.csv") if r.get("nombre") == f["nombre"]]
        if registro:
            ultimo = next((r for r in reversed(registro) if r.get("resultado") == "EXITOSO"), registro[-1])
            f["fecha"] = ultimo.get("fecha_hora", "")
            f["id_registro"] = ultimo.get("id_registro", "")
            if ultimo.get("resultado") == "EXITOSO":
                f["estado"] = "Registrada"
            else:
                f["estado"], f["detalle"] = "Error", ultimo.get("detalle", "")[:300]
        definitivo = [r for r in _leer(RESULTADOS / f"registro_definitivo_{f['lote']}.csv") if r.get("nombre") == f["nombre"]]
        if definitivo and f["estado"] == "Registrada":
            ultimo = next((r for r in reversed(definitivo) if r.get("resultado") == "EXITOSO"), definitivo[-1])
            if ultimo.get("resultado") == "EXITOSO":
                f["estado"], f["fecha"] = "Definitivo", ultimo.get("fecha_hora", f["fecha"])
            else:
                f["estado"], f["detalle"] = "Error definitivo", ultimo.get("detalle", "")[:300]
    return filas


# --- Empaquetados (hoja EMP) y promociones (hoja PROMO) ----------------------------

def _tabla(hoja, primera: str) -> tuple[list[str], list[tuple]]:
    """Cabecera y filas de la tabla cuya primera columna es «N°» y la segunda `primera`."""
    filas = list(hoja.iter_rows(values_only=True))
    inicio = next((i for i, f in enumerate(filas) if f and _texto(f[0]) == "N°"
                   and len(f) > 1 and _texto(f[1]) == primera), None)
    if inicio is None:
        return [], []
    cabecera = [_texto(c) for c in filas[inicio]]
    datos = []
    for fila in filas[inicio + 1:]:
        if not fila or not any(v not in (None, "") for v in fila):
            break
        datos.append(fila)
    return cabecera, datos


def _campo(cabecera, fila, titulo) -> str:
    return _texto(fila[cabecera.index(titulo)]) if titulo in cabecera and cabecera.index(titulo) < len(fila) else ""


def _lista(texto: str) -> list[str]:
    return [t.strip() for t in re.split(r"[;\n,]+", texto or "") if t.strip()]


def leer_paquetes() -> list[dict]:
    from openpyxl import load_workbook
    libro = load_workbook(ARCHIVO, data_only=True)
    if "EMP" not in libro.sheetnames:
        return []
    cab, filas = _tabla(libro["EMP"], "Nombre del paquete")
    salida = []
    for fila in filas:
        nombre = _campo(cab, fila, "Nombre del paquete")
        if not nombre:
            continue
        salida.append({
            "tipo": "EMP", "codigo": "EMP", "nombre": nombre, "grupo": _campo(cab, fila, "Grupo"),
            "tarifas": _lista(_campo(cab, fila, "Tarifas del paquete")),
            "servicio": _campo(cab, fila, "Servicios"), "lote": _campo(cab, fila, "Lote") or "DEMO",
            "quiere_definitivo": _campo(cab, fila, "Registro definitivo").lower().startswith("s"),
            "accion": _campo(cab, fila, "Acción"), "parametros": _campo(cab, fila, "Parámetros de la acción"),
            "esperado": _campo(cab, fila, "Resultado esperado"), "alcance": "", "problemas": [],
        })
    return salida


def leer_promos() -> list[dict]:
    from openpyxl import load_workbook
    libro = load_workbook(ARCHIVO, data_only=True)
    if "PROMO" not in libro.sheetnames:
        return []
    cab, filas = _tabla(libro["PROMO"], "Nombre de la promoción")
    salida = []
    for fila in filas:
        nombre = _campo(cab, fila, "Nombre de la promoción")
        if not nombre:
            continue
        salida.append({
            "tipo": "PROMO", "codigo": "PROMO", "nombre": nombre,
            "tarifas": _lista(_campo(cab, fila, "Tarifas base")),
            "servicio": _campo(cab, fila, "Servicios"), "lote": _campo(cab, fila, "Lote") or "DEMO",
            "quiere_definitivo": _campo(cab, fila, "Registro definitivo").lower().startswith("s"),
            "atributos_promo": _campo(cab, fila, "Atributos") or "todos",
            "variacion_porcentual": _campo(cab, fila, "Variación %") or "10",
            "variacion_cantidad": _campo(cab, fila, "Variación cantidad") or "1",
            "duracion": _campo(cab, fila, "Duración (días)") or "30",
            "accion": "", "parametros": "", "alcance": "", "problemas": [],
        })
    return salida


def estado_registro(f: dict, carpeta: str) -> None:
    """Pendiente / Registrada / Definitivo / Error, para paquetes (EMP) y promociones (PRM)."""
    f.update(estado="Pendiente", detalle="", fecha="")
    registro = [r for r in _leer(RESULTADOS / carpeta / f"resultados_{f['lote']}.csv") if r.get("nombre") == f["nombre"]]
    if registro:
        ultimo = next((r for r in reversed(registro) if r.get("resultado") == "EXITOSO"), registro[-1])
        f["fecha"] = ultimo.get("fecha_hora", "")
        if ultimo.get("resultado") == "EXITOSO":
            f["estado"] = "Registrada"
        else:
            f["estado"], f["detalle"] = "Error", (ultimo.get("detalle") or "")[:300]
    definitivo = [r for r in _leer(RESULTADOS / f"registro_definitivo_{f['lote']}.csv") if r.get("nombre") == f["nombre"]]
    if definitivo and f["estado"] == "Registrada":
        ultimo = next((r for r in reversed(definitivo) if r.get("resultado") == "EXITOSO"), definitivo[-1])
        if ultimo.get("resultado") == "EXITOSO":
            f["estado"] = "Definitivo"
        else:
            f["estado"], f["detalle"] = "Error definitivo", (ultimo.get("detalle") or "")[:300]


ACCION_SCRIPT = {"Renombre": "renombre", "Cese": "cese", "Fe de erratas": "fe_erratas",
                 "Rectificación": "rectificar", "Rectificacion": "rectificar", "Cambio": "cambio"}


def estado_acciones(filas: list[dict]) -> None:
    """Estado de la acción prevista de cada fila según resultados/ACCIONES/acciones.csv."""
    registro = _leer(RESULTADOS / "ACCIONES" / "acciones.csv")
    for f in filas:
        accion = ACCION_SCRIPT.get(f.get("accion", ""))
        f["estado_accion"], f["detalle_accion"] = ("—", "") if not accion else ("Pendiente", "")
        if not accion:
            continue
        propias = [r for r in registro if r.get("accion") == accion and r.get("nombre") == f["nombre"]]
        if propias:
            ultimo = next((r for r in reversed(propias) if r.get("resultado") == "EXITOSO"), propias[-1])
            f["estado_accion"] = {"EXITOSO": "Aplicada", "EN_RESUMEN": "Probada (resumen)",
                                  "NO_APLICA": "No aplica aún"}.get(ultimo.get("resultado"), "Error")
            f["detalle_accion"] = (ultimo.get("detalle") or "")[:300]


def todo_con_estado() -> dict:
    mono = con_estado(leer_filas())
    paquetes = leer_paquetes()
    promos = leer_promos()
    for p in paquetes:
        estado_registro(p, "EMP")
    for p in promos:
        estado_registro(p, "PRM")
    estado_acciones(mono)
    estado_acciones(paquetes)
    return {"filas": mono, "paquetes": paquetes, "promos": promos}


# --- Licitación y Contratos (datos/licitaciones.xlsx) -------------------------------

ARCHIVO_LYC = catalogo.PROYECTO / "datos" / "licitaciones.xlsx"


def leer_licitaciones() -> list[dict]:
    """Filas de datos/licitaciones.xlsx con su estado (Pendiente / Registrada / Definitivo / Error)."""
    if not ARCHIVO_LYC.exists():
        return []
    import sys
    if str(catalogo.PROYECTO) not in sys.path:
        sys.path.insert(0, str(catalogo.PROYECTO))
    from licitacion import leer_excel

    try:
        filas = leer_excel(ARCHIVO_LYC)
    except SystemExit as error:            # leer_excel avisa con sys.exit (uso en terminal)
        raise ValueError(str(error))
    salida = []
    for f in filas:
        fila = {
            "nombre": f["nombre"], "lote": f["lote"], "servicios": f["servicios"],
            "n_servicios": len(f["servicios"]), "ambito": f["ambito"],
            "valor": f["valor"], "moneda": f["moneda"], "duracion": f["duracion"],
            "quiere_definitivo": f["definitivo"].strip().lower() in ("sí", "si", "s", "x", "1", "true"),
        }
        estado_registro(fila, "LYC")
        salida.append(fila)
    return salida

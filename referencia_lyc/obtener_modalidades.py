import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright


CDP_URL = "http://localhost:9222"

ARCHIVO_ENTRADA = Path("tarifas_ejecucion_511.csv")
ARCHIVO_RESULTADOS = Path("resultados_modalidades.csv")
ARCHIVO_EXCEL = Path("codigos_modalidades.xlsx")
DIRECTORIO_CAPTURAS = Path("capturas_modalidades")

TIMEOUT = 60_000

MAPEO_SERVICIOS = {
    "Telefonía fija postpago": "TF pos",
    "Telefonía fija prepago": "TF pre",
    "Servicio público móvil postpago": "SM pos",
    "Servicio público móvil control": "SM ctrl",
    "Servicio público móvil paquete": "SM paq",
    "Televisión de paga Postpago": "TV pos",
    "Televisión de paga Prepago": "TV pre",
    "Internet fijo postpago": "IF pos",
    "Arrendamiento de circuitos": "AC",
}


def obtener_pagina_sirt(context):
    for page in context.pages:
        if "sirt.pe.be1business.com" in page.url:
            return page

    raise RuntimeError("No se encontró una pestaña abierta de SIRT.")


def leer_tarifas_objetivo():
    if not ARCHIVO_ENTRADA.exists():
        raise FileNotFoundError(f"No se encontró: {ARCHIVO_ENTRADA.resolve()}")

    with ARCHIVO_ENTRADA.open("r", newline="", encoding="utf-8-sig") as archivo:
        filas = list(csv.DictReader(archivo, delimiter=";"))

    tarifas = []
    codigos_registrados = set()

    for fila in filas:
        codigo = fila.get("codigo_sirt", "").strip()
        nombre = fila.get("nombre", "").strip()

        if not codigo or codigo in codigos_registrados:
            continue

        codigos_registrados.add(codigo)

        tarifas.append(
            {
                "codigo_sirt": codigo,
                "nombre_original": nombre,
            }
        )

    return tarifas


def leer_codigos_procesados():
    if not ARCHIVO_RESULTADOS.exists():
        return set()

    procesados = set()

    with ARCHIVO_RESULTADOS.open("r", newline="", encoding="utf-8-sig") as archivo:
        filas = csv.DictReader(archivo, delimiter=";")

        for fila in filas:
            if fila.get("resultado") != "EXITOSO":
                continue

            codigo = fila.get("codigo_sirt", "").strip()

            if codigo:
                procesados.add(codigo)

    return procesados


def registrar_resultado(codigo, version, nombre, cantidad_servicios, servicios, modalidad, descripcion, resultado, detalle):
    archivo_nuevo = not ARCHIVO_RESULTADOS.exists()

    campos = [
        "codigo_sirt",
        "version",
        "nombre",
        "cantidad_servicios",
        "servicios",
        "modalidad_resumen",
        "descripcion",
        "resultado",
        "detalle",
    ]

    with ARCHIVO_RESULTADOS.open("a", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=campos, delimiter=";")

        if archivo_nuevo:
            escritor.writeheader()

        escritor.writerow(
            {
                "codigo_sirt": codigo,
                "version": version,
                "nombre": nombre,
                "cantidad_servicios": cantidad_servicios,
                "servicios": " | ".join(servicios),
                "modalidad_resumen": modalidad,
                "descripcion": descripcion,
                "resultado": resultado,
                "detalle": detalle,
            }
        )


def abrir_tarifas_registradas(page):
    campo_codigo = page.get_by_role("textbox", name="Ingresa código tarifa")

    if campo_codigo.count() > 0 and campo_codigo.is_visible():
        return

    enlace = page.get_by_role("link", name="Tarifas registradas")

    if enlace.count() > 0 and enlace.is_visible():
        enlace.click()

    page.get_by_role("button", name="Filtro avanzado").wait_for(state="visible", timeout=TIMEOUT)


def abrir_filtro_avanzado(page):
    campo_codigo = page.get_by_role("textbox", name="Ingresa código tarifa")

    if campo_codigo.count() > 0 and campo_codigo.is_visible():
        return

    page.get_by_role("button", name="Filtro avanzado").click()
    campo_codigo.wait_for(state="visible", timeout=TIMEOUT)


def buscar_tarifa(page, codigo):
    abrir_filtro_avanzado(page)

    campo_codigo = page.get_by_role("textbox", name="Ingresa código tarifa")
    campo_codigo.fill(codigo)

    page.get_by_role("button", name="Buscar", exact=True).click()

    page.wait_for_function(
        """
        codigo => {
            const filas = Array.from(document.querySelectorAll("tbody tr"));

            return filas.some(fila => {
                const celdas = Array.from(fila.querySelectorAll("td"));
                return celdas.some(celda => celda.innerText.trim() === codigo);
            });
        }
        """,
        arg=codigo,
        timeout=TIMEOUT,
    )

    page.wait_for_timeout(500)


def convertir_version(version):
    numeros = re.findall(r"\d+", version)

    if not numeros:
        return (0,)

    return tuple(int(numero) for numero in numeros)


def obtener_fila_version_mas_alta(page, codigo):
    filas = page.locator("tbody tr")
    coincidencias = []

    for indice in range(filas.count()):
        fila = filas.nth(indice)
        celdas = fila.locator("td")

        if celdas.count() < 7:
            continue

        codigo_fila = celdas.nth(2).inner_text().strip()
        version_fila = celdas.nth(3).inner_text().strip()
        nombre_fila = celdas.nth(6).inner_text().strip()

        if codigo_fila == codigo:
            coincidencias.append(
                {
                    "fila": fila,
                    "version": version_fila,
                    "nombre": nombre_fila,
                }
            )

    if not coincidencias:
        raise RuntimeError(f"No se encontró el código {codigo} después de aplicar el filtro.")

    coincidencias.sort(key=lambda registro: convertir_version(registro["version"]), reverse=True)

    seleccionada = coincidencias[0]

    return seleccionada["fila"], seleccionada["version"], seleccionada["nombre"]


def abrir_ficha_informativa(page, fila):
    fila.get_by_label("Acciones").click()
    page.get_by_role("menuitem", name="Ficha informativa", exact=True).click()

    page.get_by_role("button", name="Salir", exact=True).wait_for(state="visible", timeout=TIMEOUT)
    page.wait_for_timeout(500)


def obtener_texto_ficha(page):
    dialogos = page.get_by_role("dialog")

    if dialogos.count() > 0:
        dialogo = dialogos.last

        if dialogo.is_visible():
            return dialogo.inner_text().strip()

    return page.locator("body").inner_text().strip()


def extraer_descripcion(texto_ficha):
    lineas = [linea.strip() for linea in texto_ficha.splitlines() if linea.strip()]

    candidatos = []

    for linea in lineas:
        linea_minuscula = linea.lower()

        if "descripción para:" in linea_minuscula:
            candidatos.append(linea)

        elif "descripción automatizada para:" in linea_minuscula:
            candidatos.append(linea)

        elif "descripción del servicio" in linea_minuscula and "|" in linea:
            candidatos.append(linea)

    if candidatos:
        return max(candidatos, key=len)

    for indice, linea in enumerate(lineas):
        if "descripción del servicio" in linea.lower() and indice + 1 < len(lineas):
            return lineas[indice + 1]

    raise RuntimeError("No se encontró la descripción en la ficha informativa.")


def identificar_servicios(descripcion):
    descripcion_normalizada = descripcion.casefold()
    servicios_encontrados = []

    for servicio in MAPEO_SERVICIOS:
        if servicio.casefold() in descripcion_normalizada:
            servicios_encontrados.append(servicio)

    return servicios_encontrados


def generar_modalidad_resumen(servicios):
    if not servicios:
        return ""

    abreviaturas = [MAPEO_SERVICIOS[servicio] for servicio in servicios]

    return " + ".join(abreviaturas)


def obtener_cantidad_servicios_nombre(nombre):
    coincidencia = re.search(r"LyC-(\d{2})S-", nombre, re.IGNORECASE)

    if not coincidencia:
        return 0

    return int(coincidencia.group(1))


def cerrar_ficha(page):
    boton_salir = page.get_by_role("button", name="Salir", exact=True)
    boton_salir.wait_for(state="visible", timeout=TIMEOUT)
    boton_salir.click()

    page.get_by_role("textbox", name="Ingresa código tarifa").wait_for(state="visible", timeout=TIMEOUT)
    page.wait_for_timeout(300)


def obtener_datos_tarifa(page, codigo):
    buscar_tarifa(page, codigo)

    fila, version, nombre = obtener_fila_version_mas_alta(page, codigo)

    abrir_ficha_informativa(page, fila)

    texto_ficha = obtener_texto_ficha(page)
    descripcion = extraer_descripcion(texto_ficha)
    servicios = identificar_servicios(descripcion)
    modalidad = generar_modalidad_resumen(servicios)

    cerrar_ficha(page)

    return version, nombre, descripcion, servicios, modalidad


def recuperar_pagina(page):
    try:
        boton_salir = page.get_by_role("button", name="Salir", exact=True)

        if boton_salir.count() > 0 and boton_salir.is_visible():
            boton_salir.click()
            page.wait_for_timeout(300)

        for _ in range(2):
            page.keyboard.press("Escape")
            page.wait_for_timeout(200)

        campo_codigo = page.get_by_role("textbox", name="Ingresa código tarifa")

        if campo_codigo.count() > 0 and campo_codigo.is_visible():
            return

        abrir_tarifas_registradas(page)
        abrir_filtro_avanzado(page)

    except Exception as error:
        print(f"No se pudo recuperar la pantalla: {error}")


def generar_excel():
    if not ARCHIVO_RESULTADOS.exists():
        raise FileNotFoundError("No existe el archivo resultados_modalidades.csv.")

    with ARCHIVO_RESULTADOS.open("r", newline="", encoding="utf-8-sig") as archivo:
        filas = [
            fila
            for fila in csv.DictReader(archivo, delimiter=";")
            if fila.get("resultado") == "EXITOSO"
        ]

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Modalidades"

    encabezados = [
        "Código SIRT",
        "Versión",
        "Nombre de tarifa",
        "Cantidad de servicios",
        "Servicios identificados",
        "Modalidad resumida",
        "Descripción de la ficha",
    ]

    hoja.append(encabezados)

    for fila in filas:
        hoja.append(
            [
                fila["codigo_sirt"],
                fila["version"],
                fila["nombre"],
                int(fila["cantidad_servicios"]),
                fila["servicios"],
                fila["modalidad_resumen"],
                fila["descripcion"],
            ]
        )

    relleno_encabezado = PatternFill(fill_type="solid", fgColor="1F4E78")

    for celda in hoja[1]:
        celda.fill = relleno_encabezado
        celda.font = Font(color="FFFFFF", bold=True)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    

    anchos = {
        1: 22,
        2: 10,
        3: 40,
        4: 22,
        5: 90,
        6: 55,
        7: 110,
    }

    for columna, ancho in anchos.items():
        hoja.column_dimensions[get_column_letter(columna)].width = ancho

    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    libro.save(ARCHIVO_EXCEL)


def ejecutar():
    DIRECTORIO_CAPTURAS.mkdir(exist_ok=True)

    tarifas = leer_tarifas_objetivo()
    codigos_procesados = leer_codigos_procesados()

    exitosas = 0
    errores = 0
    omitidas = 0

    print(f"Tarifas objetivo: {len(tarifas)}")
    print(f"Procesadas anteriormente: {len(codigos_procesados)}")

    with sync_playwright() as playwright:
        browser = playwright.chromium.connect_over_cdp(CDP_URL)
        context = browser.contexts[0]
        page = obtener_pagina_sirt(context)

        page.set_default_timeout(TIMEOUT)
        page.bring_to_front()

        abrir_tarifas_registradas(page)
        abrir_filtro_avanzado(page)

        for correlativo, tarifa in enumerate(tarifas, start=1):
            codigo = tarifa["codigo_sirt"]

            if codigo in codigos_procesados:
                omitidas += 1
                continue

            print(f"\n[{correlativo}/{len(tarifas)}] {codigo}")

            try:
                version, nombre, descripcion, servicios, modalidad = obtener_datos_tarifa(page, codigo)

                cantidad_esperada = obtener_cantidad_servicios_nombre(nombre)
                cantidad_encontrada = len(servicios)

                if cantidad_encontrada != cantidad_esperada:
                    raise RuntimeError(
                        f"Cantidad esperada según el nombre: {cantidad_esperada}. "
                        f"Cantidad identificada en la descripción: {cantidad_encontrada}. "
                        f"Descripción encontrada: {descripcion}"
                    )

                registrar_resultado(
                    codigo=codigo,
                    version=version,
                    nombre=nombre,
                    cantidad_servicios=cantidad_encontrada,
                    servicios=servicios,
                    modalidad=modalidad,
                    descripcion=descripcion,
                    resultado="EXITOSO",
                    detalle="",
                )

                codigos_procesados.add(codigo)
                exitosas += 1

                print(f"Servicios: {' | '.join(servicios)}")
                print(f"Modalidad: {modalidad}")

            except Exception as error:
                errores += 1
                captura = DIRECTORIO_CAPTURAS / f"error_{codigo}.png"

                page.screenshot(path=str(captura), full_page=True)

                registrar_resultado(
                    codigo=codigo,
                    version="",
                    nombre=tarifa["nombre_original"],
                    cantidad_servicios=0,
                    servicios=[],
                    modalidad="",
                    descripcion="",
                    resultado="ERROR",
                    detalle=str(error).replace("\n", " "),
                )

                print(f"Error: {error}")
                print(f"Captura: {captura}")

                recuperar_pagina(page)

    generar_excel()

    pendientes = len(tarifas) - len(codigos_procesados)

    print("\nPROCESO FINALIZADO")
    print("------------------")
    print(f"Procesadas correctamente: {exitosas}")
    print(f"Omitidas por ejecución anterior: {omitidas}")
    print(f"Errores: {errores}")
    print(f"Pendientes: {pendientes}")
    print(f"Resultados CSV: {ARCHIVO_RESULTADOS.resolve()}")
    print(f"Excel generado: {ARCHIVO_EXCEL.resolve()}")


if __name__ == "__main__":
    ejecutar()
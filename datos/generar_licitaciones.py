"""
Genera datos/licitaciones.xlsx: tarifas de «Licitación y Contratos» de prueba.

Una fila por tarifa, con los servicios que la forman y todos los campos que la
herramienta llena en SIRT. Por ahora, grupos de 1, 3 y 6 servicios (todas las
combinaciones de los 6 servicios elegibles):

    1 servicio  → C(6,1) =  6 tarifas
    3 servicios → C(6,3) = 20 tarifas
    6 servicios → C(6,6) =  1 tarifa
                           ---------
                           27 tarifas   (con 2, 4 y 5 serían 63 en total)

Uso (desde la carpeta del proyecto):
    python datos/generar_licitaciones.py               → grupos 1, 3 y 6
    python datos/generar_licitaciones.py --grupos 1 2 3 4 5 6

La app web y licitacion.py leen la hoja «Licitaciones».
"""

import argparse
from itertools import combinations
from pathlib import Path

PROYECTO = Path(__file__).resolve().parent.parent
SALIDA = PROYECTO / "datos" / "licitaciones.xlsx"
LOTE = "DEMO"
HOJA = "Licitaciones"

# (código corto, nombre exacto en SIRT) — de servicios.json, en el orden de SIRT.
SERVICIOS_LYC = [
    ("AC", "Arrendamiento de circuitos"),
    ("OS", "Otros servicios"),
    ("TVP", "Televisión de paga"),
    ("IF", "Internet Fijo"),
    ("SM", "Servicio Móvil"),
    ("TF", "Telefonía Fija"),
]

COLUMNAS = [
    # (título, ancho, editable)
    ("N°", 5, False),
    ("Nombre de la tarifa", 18, True),
    ("Lote", 7, True),
    ("N° servicios", 9, False),
    ("Servicios", 42, True),
    ("Duración del contrato (días)", 12, True),
    ("Identificador de contrato", 14, True),
    ("Ámbito", 9, True),
    ("Valor de contrato", 11, True),
    ("Moneda", 9, True),
    ("Documento de Empresa", 14, True),
    ("Nombre de empresa", 20, True),
    ("Descripción del Servicio", 28, True),
    ("Atributos de la tarifa", 28, True),
    ("Ámbito geográfico", 20, True),
    ("Detalle del servicio", 30, True),
    ("Registro definitivo", 10, True),
]

AMBITOS = ["Público", "Privado"]
MONEDAS = ["Soles", "Dolares", "Euro"]


def filas(grupos: list[int]) -> list[list]:
    resultado = []
    n = 0
    for tamano in grupos:
        for combo in combinations(SERVICIOS_LYC, tamano):
            n += 1
            codigos = "+".join(c for c, _ in combo)
            resultado.append([
                n,
                f"TL-LYC-{LOTE}-{n:03d}",
                LOTE,
                tamano,
                " + ".join(nombre for _, nombre in combo),
                365,
                f"CTR-{LOTE}-{n:03d}",
                AMBITOS[n % 2],
                f"{1000 + n * 50}.50",
                MONEDAS[n % 3],
                f"{20100000000 + n}",
                f"Empresa contratante de prueba {n:03d} SAC",
                f"Servicio licitado de prueba ({codigos})",
                f"Atributos de la tarifa de prueba {n:03d}",
                "Lima Metropolitana",
                "Detalle de prueba automatizada del servicio",
                "Sí",
            ])
    return resultado


def generar(grupos: list[int]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    CAB = PatternFill("solid", fgColor="1F4E78")
    EDIT = PatternFill("solid", fgColor="FFF2CC")
    LADO = Side(style="thin", color="BFBFBF")
    BORDE = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)

    datos = filas(grupos)
    wb = Workbook()
    hoja = wb.active
    hoja.title = HOJA
    hoja["A1"] = "Licitación y Contratos · tarifas de prueba"
    hoja["A1"].font = Font(name=F, bold=True, size=14)
    resumen = " · ".join(f"{g} servicio(s): {sum(1 for f in datos if f[3] == g)}" for g in grupos)
    hoja["A2"] = (f"{len(datos)} tarifas ({resumen}). Inicio de vigencia = día en que se registra; "
                  "fin = inicio + duración − 1. Celdas amarillas editables.")
    hoja["A2"].font = Font(name=F, italic=True, size=9, color="595959")

    for j, (titulo, ancho, _) in enumerate(COLUMNAS, start=1):
        c = hoja.cell(4, j, titulo)
        c.font, c.fill, c.border = Font(name=F, bold=True, color="FFFFFF", size=9), CAB, BORDE
        c.alignment = Alignment(wrap_text=True, vertical="center")
        hoja.column_dimensions[get_column_letter(j)].width = ancho
    for i, fila in enumerate(datos, start=1):
        for j, valor in enumerate(fila, start=1):
            c = hoja.cell(4 + i, j, valor)
            editable = COLUMNAS[j - 1][2]
            c.font = Font(name=F, size=9, color="0000FF" if editable else "000000")
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if editable:
                c.fill = EDIT
    hoja.freeze_panes = "C5"

    leyenda = wb.create_sheet("Leyenda")
    textos = [
        ("Licitación y Contratos", True, 14),
        ("", False, 10),
        ("Paso 1: servicios (uno por uno con «Agregar»), nombre, inicio y fin de vigencia, duración del contrato,", False, 10),
        ("        identificador de contrato, ámbito, valor y moneda, empresa y los tres textos.", False, 10),
        ("        La duración debe coincidir con las fechas: SIRT la acepta si es (fin − inicio) + 1 días.", False, 10),
        ("Paso 2: una pestaña «Servicio N» por servicio con «Detalle del servicio» → «Guardar e ir al siguiente».", False, 10),
        ("Paso 3: Confirmar. Después, registro definitivo (sin disponibilidad comercial).", False, 10),
        ("", False, 10),
        ("Servicios: separados por « + », con el nombre exacto de SIRT:", False, 10),
    ] + [(f"   {c} = {n}", False, 10) for c, n in SERVICIOS_LYC] + [
        ("Ámbito: Público / Privado · Moneda: Soles / Dolares / Euro · Valor: hasta 4 decimales.", False, 10),
        ("Documento de Empresa: solo dígitos (máx. 15).", False, 10),
    ]
    for i, (t, negrita, tam) in enumerate(textos, start=1):
        leyenda.cell(i, 1, t).font = Font(name=F, bold=negrita, size=tam)
    leyenda.column_dimensions["A"].width = 110

    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Generado: {SALIDA}  ({len(datos)} tarifas: {resumen})")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Genera datos/licitaciones.xlsx")
    parser.add_argument("--grupos", nargs="+", type=int, default=[1, 3, 6])
    generar(parser.parse_args().grupos)

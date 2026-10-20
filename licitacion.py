"""
Registra tarifas de «Licitación y Contratos» desde datos/licitaciones.xlsx.

Flujo (según grabaciones/licitacion_contratos/, 16/10):
  Paso 1: Tipo de Tarifa «Licitación y Contratos» → por cada servicio:
          «Seleccionar servicios» → opción → «Agregar».  Luego nombre, inicio y fin
          de vigencia, «Duración del contrato (días)», identificador de contrato,
          Ámbito, Valor de contrato, Moneda, Documento y Nombre de empresa,
          Descripción del Servicio, Atributos de la tarifa, Ámbito geográfico
          → «Guardar y continuar con el paso».
          La duración debe cuadrar con las fechas: SIRT aceptó 401 días para
          16/10/2026 → 20/11/2027 (diferencia 400 + 1). Si la rechaza, se prueba
          con un día menos y uno más.
  Paso 2: una pestaña «Servicio N» por servicio: «Detalle del servicio» →
          «Guardar e ir al siguiente» (en la última vuelve a la primera) →
          «Continuar con el paso».
  Paso 3: Confirmar Tarifa → Confirmar tarifa → Aceptar.

Uso:
    python licitacion.py --solo-listar
    python licitacion.py --nombres TL-LYC-DEMO-001 --diagnostico
    python licitacion.py --limite 2

Opciones:
    --excel RUTA         (por defecto datos/licitaciones.xlsx)
    --nombres N1 N2 ...  solo esas filas del Excel
    --limite N           como máximo N tarifas (las primeras pendientes)
    --repetir            vuelve a registrar aunque ya esté EXITOSA
    --solo-listar        muestra lo que haría, sin tocar SIRT
    --diagnostico        capturas por paso y mapa de la página si falla
    --al-frente          trae la ventana de Edge al frente al empezar

Resultados en resultados/LYC/resultados_<lote>.csv (mismo formato que ejecutar.py,
así registro_definitivo.py --lote <lote> los toma después).
"""

import argparse
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import comun
from ejecutar import Diagnostico
from empaquetar import abrir_nueva_tarifa, confirmar_paquete, texto_dialogo, texto_visible

CODIGO = "LYC"
TIPO_LICITACION = re.compile(r"Licitaci[oó]n", re.I)
EXCEL = Path("datos") / "licitaciones.xlsx"
HOJA = "Licitaciones"
INSTALACION_OS = "10"   # «Otros servicios» registra tarifa de instalación (si el campo aparece)

COLUMNAS = {
    "nombre": "Nombre de la tarifa",
    "lote": "Lote",
    "servicios": "Servicios",
    "duracion": "Duración del contrato (días)",
    "identificador": "Identificador de contrato",
    "ambito": "Ámbito",
    "valor": "Valor de contrato",
    "moneda": "Moneda",
    "documento": "Documento de Empresa",
    "empresa": "Nombre de empresa",
    "descripcion": "Descripción del Servicio",
    "atributos": "Atributos de la tarifa",
    "geografico": "Ámbito geográfico",
    "detalle": "Detalle del servicio",
    "definitivo": "Registro definitivo",
}


# --- Excel --------------------------------------------------------------------

SERVICIOS_LYC = ["Arrendamiento de circuitos", "Otros servicios", "Televisión de paga",
                 "Internet Fijo", "Servicio Móvil", "Telefonía Fija"]
COLUMNAS_NUEVAS = {   # automatizacion/registro_licitaciones.xlsx (cabecera en la fila 5)
    "nombre": "Nombre de la tarifa", "definitivo": "Registro definitivo", "inicio": "Inicio de vigencia",
    "duracion": "Duración del contrato (días)", "identificador": "Identificador de contrato",
    "ambito": "Ámbito", "valor": "Valor de contrato", "moneda": "Moneda",
    "documento": "Documento de Empresa", "empresa": "Nombre de empresa",
    "descripcion": "Descripción del Servicio", "atributos": "Atributos de la tarifa",
    "geografico": "Ámbito geográfico", "instalacion": "Tarifa instalación (solo con Otros servicios)",
}
LOTE_EXCEL = "EXCEL"


def leer_excel_nuevo(hoja) -> list[dict]:
    """Formato de automatizacion/registro_licitaciones.xlsx: un «incluir» y un «detalle» por servicio."""
    cabecera = [str(c.value).strip() if c.value is not None else "" for c in hoja[5]]
    faltan = [t for t in COLUMNAS_NUEVAS.values() if t not in cabecera]
    if faltan:
        sys.exit(f"Faltan columnas en la hoja {HOJA}: {', '.join(faltan)}")
    filas = []
    for valores in hoja.iter_rows(min_row=6, values_only=True):
        celda = lambda titulo: valores[cabecera.index(titulo)]
        if not valores or not celda("Nombre de la tarifa"):
            continue
        fila = {}
        for clave, titulo in COLUMNAS_NUEVAS.items():
            valor = celda(titulo)
            if clave == "inicio":
                fila[clave] = valor.date() if isinstance(valor, datetime) else valor if isinstance(valor, date) else None
            elif isinstance(valor, float) and valor.is_integer():
                fila[clave] = str(int(valor))
            else:
                fila[clave] = "" if valor is None else str(valor).strip()
        fila["servicios"], fila["detalles"] = [], {}
        for servicio in SERVICIOS_LYC:
            incluir = celda(f"{servicio}: incluir") if f"{servicio}: incluir" in cabecera else None
            if str(incluir or "").strip().lower() in ("sí", "si"):
                fila["servicios"].append(servicio)
                detalle = celda(f"{servicio}: detalle del servicio")
                fila["detalles"][servicio] = str(detalle).strip() if detalle else "Detalle del servicio"
        fila["detalle"] = "Detalle del servicio"
        fila["duracion"] = int(float(fila["duracion"] or 365))
        fila["lote"] = LOTE_EXCEL
        filas.append(fila)
    return filas


def leer_excel(ruta: Path) -> list[dict]:
    """Filas de la hoja «Licitaciones» (cabecera en la fila 4; o en la 5 con el formato nuevo)."""
    from openpyxl import load_workbook

    if not ruta.exists():
        sys.exit(f"No existe {ruta}. Genéralo con: python datos/generar_licitaciones.py")
    hoja = load_workbook(ruta, data_only=True)[HOJA]
    if any(str(c.value or "").endswith(": incluir") for c in hoja[5]):
        return leer_excel_nuevo(hoja)
    cabecera = [str(c.value).strip() if c.value is not None else "" for c in hoja[4]]
    indices = {}
    for clave, titulo in COLUMNAS.items():
        if titulo not in cabecera:
            sys.exit(f"Falta la columna «{titulo}» en la hoja {HOJA} de {ruta}.")
        indices[clave] = cabecera.index(titulo)

    filas = []
    for valores in hoja.iter_rows(min_row=5, values_only=True):
        if not valores or not valores[indices["nombre"]]:
            continue
        fila = {clave: ("" if valores[i] is None else str(valores[i]).strip()) for clave, i in indices.items()}
        fila["servicios"] = [s.strip() for s in fila["servicios"].split("+") if s.strip()]
        fila["duracion"] = int(float(fila["duracion"] or 365))
        fila["lote"] = fila["lote"] or "DEMO"
        filas.append(fila)
    return filas


# --- SIRT ---------------------------------------------------------------------

def elegir(page, combobox, opcion) -> None:
    """Lista de selección única; combobox y opción pueden ser texto exacto o regex."""
    cerrar_calendario(page)
    page.get_by_role("combobox", name=combobox).first.click()
    overlay = page.locator(".cdk-overlay-container")
    if isinstance(opcion, str):
        overlay.get_by_role("option", name=opcion, exact=True).click()
    else:
        overlay.get_by_role("option", name=opcion).first.click()


def cerrar_calendario(page) -> None:
    """Si quedó abierto un calendario (su fondo tapa los clics), lo cierra."""
    for _ in range(3):
        fondo = page.locator(".cdk-overlay-backdrop.cdk-overlay-backdrop-showing")
        if fondo.count() == 0:
            return
        page.keyboard.press("Escape")
        page.wait_for_timeout(300)


def elegir_fecha(page, etiqueta: str, fecha: date) -> None:
    """Abre el calendario y avanza mes a mes («Next month») hasta el día pedido,
    como en la grabación (escribir la fecha deja el calendario abierto)."""
    from comun import MESES

    cerrar_calendario(page)
    texto_boton = f"{fecha.day} de {MESES[fecha.month - 1]} de {fecha.year}"
    page.locator("mat-form-field").filter(has_text=etiqueta).get_by_label("Open calendar").click()
    page.get_by_role("button", name="Next month").wait_for(state="visible")
    for _ in range(60):
        boton = page.get_by_role("button", name=texto_boton, exact=True)
        if boton.count() > 0:
            boton.click()
            page.wait_for_timeout(300)
            cerrar_calendario(page)
            return
        page.get_by_role("button", name="Next month").click()
        page.wait_for_timeout(150)
    cerrar_calendario(page)
    raise RuntimeError(f"No se encontró el {texto_boton} en el calendario de «{etiqueta}».")


def agregar_servicio(page, servicio: str) -> None:
    """«Seleccionar servicios» → opción → «Agregar»."""
    page.get_by_role("combobox", name="Seleccionar servicios").click()
    opcion = page.locator(".cdk-overlay-container").get_by_role("option", name=servicio, exact=True)
    try:
        opcion.wait_for(state="visible", timeout=5_000)
    except Exception:
        page.keyboard.press("Escape")
        if texto_visible(page, servicio):
            print(f"    {servicio}: ya estaba agregado")
            return
        raise RuntimeError(f"«{servicio}» no aparece en «Seleccionar servicios».")
    opcion.click()
    page.get_by_role("button", name="Agregar", exact=True).click()
    page.wait_for_timeout(500)


def guardar_paso_1_con_duracion(page, duracion: int) -> tuple[str, int]:
    """Prueba la duración indicada y, si SIRT la rechaza, un día menos y uno más.
    Devuelve (número de registro, duración aceptada)."""
    campo = page.get_by_role("spinbutton", name=re.compile(r"Duración del contrato"))
    mensajes = []
    for intento in (duracion, duracion - 1, duracion + 1):
        campo.fill(str(intento))
        page.get_by_role("button", name="Guardar y continuar con el paso").click()
        for _ in range(60):                          # hasta ~15 s
            page.wait_for_timeout(250)
            if "step=2" in page.url:
                encontrado = re.search(r"nueva-tarifa[^/?]*/(\d+)", page.url) or re.search(r"/(\d+)\?", page.url)
                return (encontrado.group(1) if encontrado else ""), intento
            aceptar = page.get_by_role("button", name="Aceptar", exact=True)
            if aceptar.count() and aceptar.first.is_visible():
                mensajes.append(f"{intento} días: {texto_dialogo(page)[:200]}")
                aceptar.first.click()
                page.wait_for_timeout(400)
                break
        else:
            raise RuntimeError("El paso 1 no avanzó ni mostró mensaje.")
        print(f"    SIRT rechazó {intento} días; se prueba otro valor")
    raise RuntimeError("SIRT rechazó la duración del contrato: " + " | ".join(mensajes))


def completar_paso_1(page, fila: dict) -> tuple[str, int]:
    elegir(page, "Tipo de Tarifa", TIPO_LICITACION)
    page.get_by_role("combobox", name="Seleccionar servicios").wait_for(state="visible")

    for servicio in fila["servicios"]:
        agregar_servicio(page, servicio)

    inicio = fila.get("inicio") or date.today()
    fin = inicio + timedelta(days=fila["duracion"] - 1)

    page.get_by_role("textbox", name="Nombre de la tarifa").fill(fila["nombre"])
    elegir_fecha(page, "Inicio de vigencia", inicio)
    elegir_fecha(page, "Fin de vigencia", fin)
    page.get_by_role("textbox", name="Identificador de contrato").fill(fila["identificador"])
    if "Otros servicios" in fila["servicios"]:
        comun.llenar_si_editable(page.get_by_role("spinbutton", name="Tarifa Instalación"),
                                 fila.get("instalacion") or INSTALACION_OS)

    elegir(page, re.compile(r"^Ámbito"), fila["ambito"])
    page.get_by_role("textbox", name="Valor de contrato").fill(fila["valor"])
    elegir(page, re.compile(r"^Moneda"), fila["moneda"])
    page.get_by_role("textbox", name="Documento de Empresa").fill(fila["documento"])
    page.get_by_role("textbox", name="Nombre de empresa").fill(fila["empresa"])
    page.get_by_role("textbox", name="Descripción del Servicio").fill(fila["descripcion"])
    page.get_by_role("textbox", name="Atributos de la tarifa").fill(fila["atributos"])
    page.get_by_role("textbox", name="Ámbito geográfico").fill(fila["geografico"])

    print(f"    vigencia {inicio:%d/%m/%Y} → {fin:%d/%m/%Y}")
    return guardar_paso_1_con_duracion(page, fila["duracion"])


def completar_paso_2(page, fila: dict) -> None:
    """Una pestaña «Servicio N» por servicio: Detalle del servicio + Guardar."""
    total = len(fila["servicios"])
    for numero in range(1, total + 1):
        pestana = page.get_by_role("tab", name=f"Servicio {numero}", exact=True)
        if pestana.count():
            pestana.click()
            page.wait_for_timeout(500)
        elif total > 1:
            raise RuntimeError(f"No aparece la pestaña «Servicio {numero}».")

        activa = page.locator("mat-tab-body.mat-mdc-tab-body-active")
        zona = activa if activa.count() else page
        detalle = zona.get_by_role("textbox", name="Detalle del servicio")
        detalle.wait_for(state="visible")
        if fila.get("detalles"):
            # Formato nuevo: un detalle por servicio; la pestaña se reconoce por el servicio que muestra.
            texto_pestana = zona.inner_text().lower()
            servicio = next((s for s in fila["servicios"] if s.lower() in texto_pestana), fila["servicios"][numero - 1])
            detalle.fill(fila["detalles"][servicio])
        else:
            detalle.fill(f"{fila['detalle']} ({numero}/{total})")
        guardar = zona.get_by_role("button", name=re.compile(r"^Guardar( e ir al siguiente)?$"))
        guardar.first.click()
        page.wait_for_timeout(1_000)

        aceptar = page.get_by_role("button", name="Aceptar", exact=True)
        if aceptar.count() and aceptar.first.is_visible():
            mensaje = texto_dialogo(page)
            aceptar.first.click()
            if re.search(r"error|obligatori|complet|valid", mensaje, re.I):
                raise RuntimeError(f"Servicio {numero}: {mensaje[:300]}")

    page.get_by_role("button", name="Continuar con el paso").click()
    confirmar = page.get_by_role("button", name="Confirmar Tarifa", exact=True)
    try:
        confirmar.wait_for(state="visible", timeout=15_000)
    except Exception:
        raise RuntimeError(f"No pasó al resumen. Mensaje de SIRT: {texto_dialogo(page)[:300]}")


# --- Principal ----------------------------------------------------------------

def leer_argumentos():
    parser = argparse.ArgumentParser(description="Registra tarifas de Licitación y Contratos.")
    parser.add_argument("--excel", default=str(EXCEL))
    parser.add_argument("--nombres", nargs="+")
    parser.add_argument("--limite", type=int, default=0)
    parser.add_argument("--repetir", action="store_true")
    parser.add_argument("--solo-listar", action="store_true")
    parser.add_argument("--diagnostico", action="store_true")
    parser.add_argument("--al-frente", action="store_true")
    return parser.parse_args()


def exitosas(lote: str) -> set[str]:
    import csv
    ruta = comun.rutas_lote(CODIGO, lote)["resultados"]
    if not ruta.exists():
        return set()
    with ruta.open("r", newline="", encoding="utf-8-sig") as entrada:
        return {f["nombre"] for f in csv.DictReader(entrada, delimiter=";") if f.get("resultado") == "EXITOSO"}


def main() -> None:
    args = leer_argumentos()
    filas = leer_excel(Path(args.excel))

    if args.nombres:
        pedidos = set(args.nombres)
        faltan = pedidos - {f["nombre"] for f in filas}
        if faltan:
            print("No están en el Excel: " + ", ".join(sorted(faltan)))
        filas = [f for f in filas if f["nombre"] in pedidos]

    pendientes = []
    for fila in filas:
        if not args.repetir and fila["nombre"] in exitosas(fila["lote"]):
            print(f"  {fila['nombre']}: ya registrada (EXITOSO), se omite")
            continue
        pendientes.append(fila)
    if args.limite:
        pendientes = pendientes[: args.limite]

    print(f"\nTarifas de Licitación y Contratos a registrar: {len(pendientes)}")
    for f in pendientes:
        print(f"  - {f['nombre']:<18} {len(f['servicios'])} serv.: {' + '.join(f['servicios'])}")
    if args.solo_listar or not pendientes:
        if args.solo_listar:
            print("\n--solo-listar: no se tocó SIRT.")
        return

    from playwright.sync_api import sync_playwright

    errores = 0
    with sync_playwright() as playwright:
        browser = playwright.chromium.connect_over_cdp(comun.CDP_URL)
        context = browser.contexts[0]
        page = comun.obtener_pagina_sirt(context)
        page.set_default_timeout(comun.TIMEOUT)
        if args.al_frente:
            page.bring_to_front()

        for indice, fila_excel in enumerate(pendientes, start=1):
            rutas = comun.rutas_lote(CODIGO, fila_excel["lote"])
            diag = Diagnostico(args.diagnostico, False, rutas["base"] / f"diagnostico_{fila_excel['lote']}")
            diag.iniciar(context)
            correlativo = int(re.search(r"(\d+)$", fila_excel["nombre"]).group(1)) if re.search(r"\d+$", fila_excel["nombre"]) else indice
            diag.iniciar_caso({"correlativo": correlativo, "nombre": fila_excel["nombre"]})
            estado = {"etapa": "", "id_registro": ""}
            resultado = {
                "correlativo": correlativo, "nombre": fila_excel["nombre"],
                "clave": " + ".join(fila_excel["servicios"]),
                "id_registro": "", "resultado": "", "detalle": "", "captura": "",
                "fecha_hora": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }
            print(f"\n[{indice}/{len(pendientes)}] {fila_excel['nombre']} ({len(fila_excel['servicios'])} servicios)")

            def etapa(nombre, accion):
                estado["etapa"] = nombre
                valor = accion()
                diag.captura_paso(page, nombre)
                return valor

            try:
                etapa("abrir_nueva_tarifa", lambda: abrir_nueva_tarifa(page))
                id_registro, dias = etapa("paso_1_guardado", lambda: completar_paso_1(page, fila_excel))
                estado["id_registro"] = id_registro
                print(f"    paso 1 guardado (registro {id_registro}, duración {dias} días)")
                etapa("paso_2_completo", lambda: completar_paso_2(page, fila_excel))
                etapa("confirmada", lambda: confirmar_paquete(page))
                resultado["resultado"] = "EXITOSO"
                resultado["detalle"] = f"duración {dias} días"
                print("    CONFIRMADA")
            except Exception as error:
                errores += 1
                rutas["capturas"].mkdir(parents=True, exist_ok=True)
                captura = rutas["capturas"] / f"{correlativo:03d}_error.png"
                try:
                    page.screenshot(path=str(captura), full_page=True)
                except Exception:
                    pass
                diag.registrar_fallo(page, error, estado["etapa"])
                resultado["resultado"] = "ERROR"
                resultado["detalle"] = f"[{estado['etapa']}] " + str(error).replace("\n", " ")[:500]
                resultado["captura"] = str(captura)
                print(f"    ERROR en '{estado['etapa']}': {error}")

            diag.cerrar_caso(resultado["resultado"] == "ERROR")
            diag.finalizar()
            resultado["id_registro"] = estado["id_registro"]
            comun.registrar_resultado(rutas["resultados"], resultado)

    print(f"\nTerminado: {len(pendientes) - errores} exitosas, {errores} con error.")
    print(f"Resultados: {comun.rutas_lote(CODIGO, pendientes[0]['lote'])['resultados'].resolve()}")
    if errores:
        sys.exit(1)


if __name__ == "__main__":
    main()

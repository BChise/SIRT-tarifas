"""
Piezas comunes de los comandos que trabajan con los Excel de automatizacion/.

- Lectura de una hoja con cabecera en la fila 5 (formato de todas las plantillas).
- Escritura del resultado en las columnas grises (Estado, Código SIRT, Fecha, Detalle).
- Ejecución de los scripts ya probados (ejecutar, empaquetar, promocionar…) como
  subproceso, mostrando su salida en pantalla.
- Registro definitivo de una tarifa (registro_definitivo.py) y lectura de su resultado.
"""

import csv
import os
import shutil
import subprocess
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

PROYECTO = Path(__file__).resolve().parent
AUTOMATIZACION = PROYECTO / "automatizacion"
RESULTADOS = PROYECTO / "resultados"
FILA_CABECERA = 5
LOTE_EXCEL = "EXCEL"          # lote con que se registra todo lo que viene de los Excel

COLORES = {
    "Registrada": "C6EFCE", "Aplicada": "C6EFCE", "Registro definitivo": "A9D08E",
    "Error": "FFC7CE", "Error de datos": "FFC7CE", "No aplica": "FFEB9C",
    "Ya registrada": "DDEBF7", "Ya aplicada": "DDEBF7", "Omitida": "EDEDED",
    "Hasta resumen": "FFF2CC",
}
SI = ("sí", "si")
INICIO = time.monotonic()      # para el «Tiempo total» del resumen


def tiempo_total(cantidad: int = 0) -> str:
    """«1 h 10 min 5 s» desde que empezó el comando (y el promedio si se da la cantidad)."""
    segundos = int(time.monotonic() - INICIO)
    h, resto = divmod(segundos, 3600)
    m, s = divmod(resto, 60)
    texto_total = (f"{h} h " if h else "") + (f"{m} min " if h or m else "") + f"{s} s"
    if cantidad:
        texto_total += f" · promedio por fila: {segundos / cantidad:.0f} s"
    return texto_total


class DatoInvalido(ValueError):
    """Un dato del Excel que impide ejecutar la fila (se escribe como «Error de datos»)."""


# --- Valores ------------------------------------------------------------------

def texto(valor) -> str:
    """Celda → texto (50.0 → «50»; None → «»)."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else repr(valor)
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor).strip()


def es_si(valor) -> bool:
    return texto(valor).lower() in SI


def lista(valor) -> list[str]:
    """«A + B + C» → [A, B, C]. También acepta «;»."""
    t = texto(valor)
    if not t:
        return []
    separador = ";" if ";" in t and "+" not in t else "+"
    return [p.strip() for p in t.split(separador) if p.strip()]


def leer_fecha(valor, campo: str = "fecha") -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    t = texto(valor)
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(t, formato).date()
        except ValueError:
            pass
    raise DatoInvalido(f"«{campo}»: fecha «{t}» no reconocida (usa dd/mm/aaaa)")


def numero(valor, campo: str, minimo=None, maximo=None, entero=False) -> str:
    """Valida un número del Excel y lo devuelve como texto para SIRT."""
    t = texto(valor).replace(",", ".")
    try:
        n = float(t)
    except ValueError:
        raise DatoInvalido(f"«{campo}» debe ser un número (tiene «{texto(valor)}»)")
    if entero and not n.is_integer():
        raise DatoInvalido(f"«{campo}» debe ser un número entero (tiene «{t}»)")
    if minimo is not None and n < minimo:
        raise DatoInvalido(f"«{campo}» debe ser al menos {minimo}")
    if maximo is not None and n > maximo:
        raise DatoInvalido(f"«{campo}» no puede pasar de {maximo}")
    return str(int(n)) if n.is_integer() else t


# --- Excel ----------------------------------------------------------------------

def excel_esta_abierto(ruta: Path) -> bool:
    """Windows: Excel bloquea el archivo. Linux/LibreOffice: deja un archivo .~lock."""
    if (ruta.parent / f".~lock.{ruta.name}#").exists() or (ruta.parent / f"~${ruta.name}").exists():
        return True
    try:
        with ruta.open("r+b"):
            return False
    except PermissionError:
        return True


def abrir_libro(ruta: Path, solo_listar: bool):
    import warnings
    from openpyxl import load_workbook
    if not ruta.exists():
        sys.exit(f"No existe {ruta}")
    if not solo_listar and excel_esta_abierto(ruta):
        sys.exit(f"El Excel está abierto: ciérralo y vuelve a correr el comando.\n  {ruta}")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")          # «Data Validation extension is not supported»
        libro = load_workbook(ruta)
    restaurar_validaciones(libro, ruta)
    return libro


def restaurar_validaciones(libro, ruta: Path) -> int:
    """Al guardar en Excel, las listas desplegables que apuntan a otra hoja («Listas») quedan
    como validaciones «x14», que openpyxl no lee y borraría al guardar. Se vuelven a crear
    como validaciones normales (Excel las entiende igual). Devuelve cuántas se restauraron."""
    import re
    import zipfile
    from openpyxl.worksheet.datavalidation import DataValidation
    try:
        with zipfile.ZipFile(ruta) as z:
            libro_xml = z.read("xl/workbook.xml").decode("utf-8")
            rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            destinos = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
            destinos.update({i: t for t, i in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
            total = 0
            for nombre, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', libro_xml):
                nombre = nombre.replace("&amp;", "&")
                destino = destinos.get(rid, "")
                ruta_xml = destino.lstrip("/") if destino.startswith("/") else "xl/" + destino
                if nombre not in libro.sheetnames or ruta_xml not in z.namelist():
                    continue
                xml = z.read(ruta_xml).decode("utf-8")
                for bloque in re.findall(r"<x14:dataValidation\b(.*?)</x14:dataValidation>", xml, re.S):
                    atributos = dict(re.findall(r'(\w+)="([^"]*)"', bloque.split(">", 1)[0]))
                    formula = re.search(r"<x14:formula1>\s*<xm:f>(.*?)</xm:f>", bloque, re.S)
                    rango = re.search(r"<xm:sqref>(.*?)</xm:sqref>", bloque, re.S)
                    if not formula or not rango:
                        continue
                    dv = DataValidation(type=atributos.get("type", "list"),
                                        formula1=formula.group(1).replace("&amp;", "&"),
                                        allow_blank=atributos.get("allowBlank") == "1")
                    dv.showErrorMessage = atributos.get("showErrorMessage") == "1"
                    if atributos.get("error"):
                        dv.error = atributos["error"].replace("&amp;", "&")
                    dv.sqref = rango.group(1)
                    libro[nombre].add_data_validation(dv)
                    total += 1
            return total
    except Exception as error:
        print(f"Aviso: no se pudieron revisar las listas desplegables del Excel ({error}).")
        return 0


COLUMNA_ID = "ID de registro"


def codigo_sirt(nombre: str, con_id: bool = False):
    """Código SIRT (p. ej. TERDS2026000123) anotado al hacer el registro definitivo
    (con_id=True → (código, ID interno))."""
    import comun
    return comun.codigo_sirt_registrado(nombre, con_id)


class Hoja:
    """Una hoja con cabecera en la fila 5 y columnas de resultado al final."""

    def __init__(self, hoja, obligatorias=(), estado="Estado", codigo="Código SIRT",
                 fecha=("Fecha de registro", "Fecha de ejecución"), detalle="Detalle"):
        self.hoja = hoja
        self.col = {}
        for celda in hoja[FILA_CABECERA]:
            if celda.value is not None and texto(celda.value) not in self.col:
                self.col[texto(celda.value)] = celda.column
        faltan = [t for t in obligatorias if t not in self.col]
        if faltan:
            sys.exit(f"Faltan columnas en la hoja «{hoja.title}»: {', '.join(faltan)}")
        fechas = (fecha,) if isinstance(fecha, str) else fecha
        self.c_estado = self.col.get(estado)
        self.c_codigo = self.col.get(codigo)
        self.c_fecha = next((self.col[f] for f in fechas if f in self.col), None)
        self.c_detalle = self.col.get(detalle)
        self.c_nombre = next((c for t, c in self.col.items() if t.startswith("Nombre")), None)
        self.c_id = self.col.get(COLUMNA_ID)

    def _columna_id(self) -> int:
        """Crea (si falta) la columna «ID de registro» al final, con el estilo de «Detalle»."""
        if self.c_id:
            return self.c_id
        from copy import copy
        from openpyxl.utils import get_column_letter
        self.c_id = max(self.col.values()) + 1
        modelo = self.hoja.cell(FILA_CABECERA, self.c_detalle or self.c_estado or 1)
        celda = self.hoja.cell(FILA_CABECERA, self.c_id, COLUMNA_ID)
        for atributo in ("font", "fill", "border", "alignment"):
            setattr(celda, atributo, copy(getattr(modelo, atributo)))
        self.hoja.column_dimensions[get_column_letter(self.c_id)].width = 13
        self.col[COLUMNA_ID] = self.c_id
        return self.c_id

    def filas(self, columna_nombre: str):
        """(número de fila, {título: valor}) de las filas con nombre."""
        c_nombre = self.col[columna_nombre]
        for r in range(FILA_CABECERA + 1, self.hoja.max_row + 1):
            if texto(self.hoja.cell(r, c_nombre).value):
                yield r, {t: self.hoja.cell(r, c).value for t, c in self.col.items()}

    def escribir(self, r: int, estado: str, codigo: str = "", fecha: str = "", detalle: str = "") -> None:
        """«codigo» puede ser el Código SIRT o el ID interno (solo dígitos): el ID va a la columna
        «ID de registro» y en «Código SIRT» queda el código real si ya se conoce."""
        from openpyxl.styles import Font, PatternFill
        codigo = texto(codigo)
        actual = texto(self.hoja.cell(r, self.c_codigo).value) if self.c_codigo else ""
        identificador = codigo if codigo.isdigit() else (actual if actual.isdigit() else "")
        nombre = texto(self.hoja.cell(r, self.c_nombre).value) if self.c_nombre else ""
        conocido, id_conocido = codigo_sirt(nombre, True) if nombre else ("", "")
        identificador = identificador or id_conocido
        if identificador:
            celda = self.hoja.cell(r, self._columna_id(), identificador)
            celda.font = Font(name="Arial", size=9)
        if not codigo or codigo.isdigit():
            codigo = conocido
            if not codigo and actual and not actual.isdigit():
                codigo = actual                                     # no borrar un código ya escrito
        for c, valor in ((self.c_estado, estado), (self.c_codigo, codigo),
                         (self.c_fecha, fecha), (self.c_detalle, detalle[:1000] if detalle else detalle)):
            if c:
                celda = self.hoja.cell(r, c, valor)
                celda.font = Font(name="Arial", size=9, bold=c == self.c_estado)
        if self.c_estado:
            self.hoja.cell(r, self.c_estado).fill = PatternFill("solid", fgColor=COLORES.get(estado, "EDEDED"))

    def estado(self, r: int) -> str:
        return texto(self.hoja.cell(r, self.c_estado).value) if self.c_estado else ""


def guardar(libro, ruta: Path) -> None:
    try:
        libro.save(ruta)
    except PermissionError:
        print("   Aviso: no se pudo guardar el Excel (¿se abrió?); se guardará al final en una copia.")


def guardar_con_copia(libro, ruta: Path, carpeta: str, prefijo: str) -> Path:
    copia = RESULTADOS / carpeta / f"{prefijo}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    copia.parent.mkdir(parents=True, exist_ok=True)
    libro.save(copia)                      # primero la copia: pase lo que pase, no se pierde nada
    guardar_final(libro, ruta, copia)
    return copia


def guardar_final(libro, ruta: Path, copia: Path) -> bool:
    """Guarda en el Excel original. Si está abierto, pide cerrarlo y escribir 'si' (insiste
    mientras siga abierto). 'copia' (o consola sin respuesta) = dejarlo solo en la copia."""
    while True:
        try:
            libro.save(ruta)
            return True
        except PermissionError:
            print("\n" + "!" * 70)
            print(f"El Excel está ABIERTO y no se puede guardar en él:\n  {ruta}")
            print(f"Los resultados ya están a salvo en la copia:\n  {copia}")
            print("Ciérralo y escribe 'si' para guardarlos también en el original"
                  " (o 'copia' para dejarlos solo en la copia).")
            try:
                respuesta = input("> ").strip().lower()
            except EOFError:                   # sin consola (p. ej. tarea programada)
                respuesta = "copia"
            if respuesta == "copia":
                print("Resultados solo en la copia.")
                return False
            if respuesta not in SI:
                print("Respuesta no reconocida.")


def confirmar(mensaje: str, si: bool) -> bool:
    if si:
        return True
    print(f"\nATENCIÓN: {mensaje}")
    return input("Escribe 'si' para continuar: ").strip().lower() in SI


def ejecutar_fila(fila: dict) -> str:
    """«Ejecutar»: vacío = Sí. Devuelve 'si', 'no' o el valor inválido."""
    t = texto(fila.get("Ejecutar")).lower()
    if t in ("", "sí", "si"):
        return "si"
    if t == "no":
        return "no"
    return t


# --- Subprocesos y CSV --------------------------------------------------------------

def correr(argumentos: list[str]) -> str:
    """Corre un script del proyecto, muestra su salida (con sangría) y la devuelve."""
    comando = [sys.executable, "-u", *argumentos]
    entorno = {**os.environ, "PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8"}   # tildes y «» en Windows
    proceso = subprocess.Popen(comando, cwd=PROYECTO, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                               text=True, encoding="utf-8", errors="replace", env=entorno)
    lineas = []
    for linea in proceso.stdout:
        lineas.append(linea)
        if linea.strip() and not linea.startswith("=" * 20):
            print("   " + linea.rstrip())
    proceso.wait()
    salida = "".join(lineas)
    puerto = os.environ.get("SIRT_PUERTO", "9222")
    if "ECONNREFUSED" in salida and puerto in salida:
        sys.exit(f"\nNo se pudo conectar con Edge (puerto {puerto}). Ábrelo en modo controlable, inicia sesión "
                 "en SIRT y vuelve a correr el comando (lo ya hecho quedó anotado en el Excel).")
    return salida


def marca_tiempo() -> str:
    return (datetime.now() - timedelta(seconds=2)).strftime("%Y-%m-%d %H:%M:%S")


def ultimo_resultado(ruta: Path, nombre: str, desde: str, clave_nombre: str = "nombre") -> dict | None:
    if not ruta.exists():
        return None
    with ruta.open("r", newline="", encoding="utf-8-sig") as entrada:
        filas = [f for f in csv.DictReader(entrada, delimiter=";")
                 if f.get(clave_nombre) == nombre and (f.get("fecha_hora") or "") >= desde]
    return filas[-1] if filas else None


def exitosa_en(ruta: Path, nombre: str) -> dict | None:
    if not ruta.exists():
        return None
    with ruta.open("r", newline="", encoding="utf-8-sig") as entrada:
        filas = [f for f in csv.DictReader(entrada, delimiter=";")
                 if f.get("nombre") == nombre and f.get("resultado") == "EXITOSO"]
    return filas[-1] if filas else None


def resumen_salida(salida: str, largo: int = 400) -> str:
    return " | ".join(l.strip() for l in salida.splitlines() if l.strip())[-largo:]


DEFINITIVOS_PRECARGADOS: dict[str, tuple[bool, str]] = {}


def precargar_definitivos(pedidos: list[tuple[str, str]], lote: str, tamano: int = 10,
                          diagnostico: bool = False) -> None:
    """Modo EN BLOQUE: hace de una vez el registro definitivo de [(nombre, disponibilidad)]
    (una llamada a registro_definitivo.py --en-bloque por cada disponibilidad distinta). Luego
    hacer_registro_definitivo() devuelve el resultado de cada una sin volver a SIRT."""
    ruta = RESULTADOS / f"registro_definitivo_{lote}.csv"
    pendientes = [(n, u) for n, u in pedidos if not exitosa_en(ruta, n)]
    if not pendientes:
        return
    grupos: dict[str, list[str]] = {}
    for nombre, ubigeos in pendientes:
        grupos.setdefault(ubigeos or "", []).append(nombre)
    for ubigeos, nombres in grupos.items():
        desde = marca_tiempo()
        argumentos = ["registro_definitivo.py", "--lote", lote, "--si", "--en-bloque", "--tamano-bloque", str(tamano)]
        for nombre in nombres:
            argumentos += ["--nombre", nombre]
        if ubigeos:
            argumentos += ["--ubigeos", ubigeos]
        if diagnostico:
            argumentos.append("--diagnostico")
        print(f"   → registro definitivo EN BLOQUE de {len(nombres)} tarifa(s)"
              + (f" (disponibilidad: {ubigeos})" if ubigeos else ""))
        salida = correr(argumentos)
        for nombre in nombres:
            fila = ultimo_resultado(ruta, nombre, desde)
            if fila is None:
                DEFINITIVOS_PRECARGADOS[nombre] = (False, "registro definitivo sin resultado: " + resumen_salida(salida))
            elif fila.get("resultado") == "EXITOSO":
                disp = fila.get("disponibilidad", "")
                modo = " en bloque" if fila.get("detalle", "").startswith("en bloque") else ""
                DEFINITIVOS_PRECARGADOS[nombre] = (True, f"registro definitivo{modo} OK" + (f" (disponibilidad: {disp})" if disp else ""))
            else:
                DEFINITIVOS_PRECARGADOS[nombre] = (False, "registro definitivo con error: " + fila.get("detalle", ""))


def hacer_registro_definitivo(nombre: str, lote: str, ubigeos: str = "", diagnostico: bool = False) -> tuple[bool, str]:
    """Registro definitivo (con disponibilidad si el servicio la pide) de UNA tarifa.
    Devuelve (éxito, detalle)."""
    if nombre in DEFINITIVOS_PRECARGADOS:
        return DEFINITIVOS_PRECARGADOS.pop(nombre)
    ruta = RESULTADOS / f"registro_definitivo_{lote}.csv"
    if exitosa_en(ruta, nombre):
        return True, "ya tenía registro definitivo"
    desde = marca_tiempo()
    argumentos = ["registro_definitivo.py", "--lote", lote, "--nombre", nombre, "--si"]
    if ubigeos:
        argumentos += ["--ubigeos", ubigeos]
    if diagnostico:
        argumentos.append("--diagnostico")
    print(f"   → registro definitivo de {nombre}" + (f" (disponibilidad: {ubigeos})" if ubigeos else ""))
    salida = correr(argumentos)
    fila = ultimo_resultado(ruta, nombre, desde)
    if fila is None:
        return False, "registro definitivo sin resultado: " + resumen_salida(salida)
    if fila.get("resultado") == "EXITOSO":
        disp = fila.get("disponibilidad", "")
        return True, "registro definitivo OK" + (f" (disponibilidad: {disp})" if disp else "")
    return False, "registro definitivo con error: " + fila.get("detalle", "")

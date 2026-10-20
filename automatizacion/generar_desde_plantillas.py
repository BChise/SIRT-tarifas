"""
Genera automatizacion/registro_tarifas_establecidas.xlsx (BORRADOR, sin conectar con
los scripts) a partir de las plantillas de carga de SIRT con tarifas de muestra:
grabaciones/Creación/<carpeta>/*(C).xlsx  (13 servicios; AC y OS no vienen ahí y se
arman con los valores de la herramienta).

Qué hace con cada plantilla:
- Hoja «Datos» → una hoja por servicio, una fila por tarifa.
- Códigos de SIRT (001001, 010001,010002, 024005…) → sus nombres (hoja «Maestras»),
  varios separados por « + »; «Seleccione una opción…» → vacío.
- Minutos/SMS («¿Es ilimitado?» + cantidad + tarifa) → 3 columnas «Atributo: …».
- Hojas de tablas (países, velocidades, bolsas) → 3 hojas «Tablas – …», cada fila
  unida a su tarifa por el NOMBRE (no por el N°).

Uso (desde la carpeta del proyecto):
    python automatizacion/generar_desde_plantillas.py
"""

import glob
import importlib
import json
import re
import sys
from pathlib import Path

PROYECTO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROYECTO))
sys.path.insert(0, str(PROYECTO / "datos"))

import comun  # noqa: E402
from app import catalogo  # noqa: E402

ORIGEN = PROYECTO / "grabaciones" / "Creación"
SALIDA = PROYECTO / "automatizacion" / "registro_tarifas_establecidas.xlsx"
FILAS_CON_LISTAS = 1000
DISPONIBILIDAD_POR_DEFECTO = "Lima"

# Carpeta de la plantilla → código de la herramienta.
CARPETAS = {
    "TF POS": "TF", "TF PRE": "TFP", "TF OTROS": "OTF", "TF PQT": "PTF",
    "IF POS": "IF", "IF OTROS": "PIF",
    "TV POS": "TVP", "TV PRE": "TVPR", "TV OTROS": "PTV",
    "SM POS": "SMP", "SM CONTROL": "SMC", "SM OTROS": "SMO", "SM PQT": "SMQ",
}
ORDEN = ["TF", "TFP", "OTF", "PTF", "IF", "PIF", "TVP", "TVPR", "PTV", "SMP", "SMC", "SMO", "SMQ", "OS", "AC"]
CON_DISPONIBILIDAD = {"TF", "TFP", "OTF", "PTF", "IF", "PIF", "TVP", "TVPR", "PTV"}

# Tipos de tabla según las columnas de su hoja en la plantilla.
TIPOS_TABLA = {
    "Países": ["Tipo", "Tarifa", "Países"],
    "Velocidades": ["Tecnología / Tipo", "Descarga máx. contratada (Mbps)", "Descarga mín. garantizada (Mbps)",
                    "Carga máx. contratada (Mbps)", "Carga mín. garantizada (Mbps)"],
    "Bolsas": ["Tipo", "Gigas", "Plataformas"],
}

# Paso 1: obligatoriedad (no está en el JSON del paso 2). Vacío en fechas = día del registro.
PASO_1 = {"Alcance de la tarifa": True, "Inicio de vigencia": True, "Inicio comercialización": True,
          "Tarifa instalación": True, "Enlace Web": True, "Fono Atención al cliente": True,
          "Email de atención al cliente": True, "Otros datos generales": False}
# 19/10 (captura de IF): el paso 1 marca «Este apartado es obligatorio» en nombre, alcance,
# inicio de vigencia, inicio de comercialización, tarifa de instalación, enlace web, fono
# y email. Fin de vigencia / comercialización y valor del plan vienen bloqueados.
SIN_JSON = {"Periodicidad": True, "Tipo de Acceso (IF)": True, "Tecnología (IF)": True,
            "Internet Ilimitado": True}
# Correcciones vistas en pantalla (el JSON de IF se toma de PIF y ahí no es obligatoria).
OBLIGATORIOS_POR_SERVICIO = {"IF": {"Valor del servicio", "Periodicidad", "Tipo de Acceso (IF)", "Tecnología (IF)"}}
SELECCIONE = re.compile(r"^seleccione una opci", re.I)


def tipo_de_tabla(cabecera: tuple) -> str:
    texto = " ".join(str(c or "") for c in cabecera).lower()
    if "gigas" in texto:
        return "Bolsas"
    if "velocidad" in texto:
        return "Velocidades"
    return "Países"


def maestras(wb) -> tuple[dict, dict]:
    codigos, paises = {}, {}
    if "Maestras" not in wb.sheetnames:
        return codigos, paises
    for tipo, codigo, descripcion in wb["Maestras"].iter_rows(min_row=2, max_col=3, values_only=True):
        if codigo is None:
            continue
        if str(tipo).upper() == "PAIS":
            paises[str(codigo).strip()] = str(descripcion).strip()
        else:
            codigos[str(codigo).strip()] = str(descripcion).strip()
    return codigos, paises


def decodificar(valor, codigos: dict, paises: dict | None = None):
    if valor is None:
        return ""
    texto = str(valor).strip()
    if SELECCIONE.match(texto):
        return ""
    if paises is not None and re.fullmatch(r"\d+(,\d+)*", texto):
        return " + ".join(paises.get(p, p) for p in texto.split(","))
    if re.fullmatch(r"\d{6}(,\d{6})*", texto) and texto.split(",")[0] in codigos:
        return " + ".join(codigos.get(c, c) for c in texto.split(","))
    return valor if isinstance(valor, (int, float)) else texto


def columnas_unicas(cabecera: list) -> list[str]:
    """«Minutos X ¿Es ilimitado?», «Cantidad…», «Tarifa…» → «Minutos X: ¿Es ilimitado?», «Minutos X: Cantidad…»."""
    salida, grupo, restantes = [], None, 0
    for titulo in cabecera:
        titulo = (titulo or "").strip()
        if titulo.endswith("¿Es ilimitado?"):
            grupo = titulo.replace("¿Es ilimitado?", "").strip()
            salida.append(f"{grupo}: ¿Es ilimitado?")
            restantes = 2
        elif restantes and grupo:
            salida.append(f"{grupo}: {titulo}")
            restantes -= 1
        else:
            salida.append(titulo)
            grupo, restantes = None, 0
    return salida


def obligatoriedad_json(archivo: str) -> dict:
    ruta = PROYECTO / "grabaciones" / archivo / "paso2.json"
    if not ruta.exists() or ruta.stat().st_size == 0:
        return {}
    datos = json.loads(ruta.read_text(encoding="utf-8"))["data"]
    salida = {}
    for seccion in datos.get("secciones") or []:
        for apartado in seccion.get("apartados") or []:
            salida[catalogo._sin_tildes(apartado["nombre"])] = bool(apartado.get("esObligatorio"))
    return salida


def es_obligatorio(columna: str, tabla: dict, codigo: str = "") -> bool:
    if columna in OBLIGATORIOS_POR_SERVICIO.get(codigo, ()):
        return True
    if codigo in VELOCIDADES_EN_LINEA and ":" in columna and columna.split(":")[0] in VELOCIDADES_EN_LINEA[codigo]:
        return (codigo, columna.split(":")[0]) in OBLIGATORIAS_EN_LINEA
    if columna in PASO_1:
        return PASO_1[columna]
    if ":" in columna and not columna.endswith("¿Es ilimitado?") and \
            re.search(r": (Cantidad|Tarifa)", columna) and "Velocidad" not in columna:
        return False                     # condicional: solo si «¿Es ilimitado?» = No
    if columna.endswith(": archivo Excel"):
        columna = columna.replace(": archivo Excel", "")
    base = columna.split(":")[0].strip()
    clave = catalogo._sin_tildes(base)
    if clave in tabla:
        return tabla[clave]
    return SIN_JSON.get(base, True)


def leer_plantilla(ruta: Path, codigo: str) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(ruta, data_only=True)
    codigos, paises = maestras(wb)
    filas = list(wb["Datos"].iter_rows(values_only=True))
    i_cab = next(i for i, f in enumerate(filas) if f and f[0] == "N°")
    secciones_crudas = filas[i_cab - 1] if i_cab else ()
    ancho = max(k for k, c in enumerate(filas[i_cab]) if c) + 1
    cabecera = columnas_unicas(list(filas[i_cab][:ancho]))
    seccion, secciones = "", []
    for k in range(ancho):
        if k < len(secciones_crudas) and secciones_crudas[k]:
            seccion = str(secciones_crudas[k]).strip()
        secciones.append(seccion)
    tarifas, nombre_por_n = [], {}
    for fila in filas[i_cab + 1:]:
        if not fila or fila[1] in (None, ""):
            continue
        valores = [decodificar(v, codigos) for v in fila[:ancho]]
        n = int(float(fila[0])) if fila[0] not in (None, "") else len(tarifas) + 1
        nombre = f"TE-{codigo}-CM-{len(tarifas) + 1:03d}"
        nombre_por_n[n] = nombre
        tarifas.append({"nombre": nombre, "original": str(fila[1]).strip(),
                        "valores": dict(zip(cabecera, valores))})
    tablas = []
    for hoja in wb.sheetnames:
        if hoja in ("Datos", "Version", "Listas", "Maestras"):
            continue
        ws = wb[hoja]
        tipo = tipo_de_tabla(next(ws.iter_rows(max_row=1, values_only=True)))
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if not fila or fila[0] in (None, ""):
                continue
            n = int(float(fila[0]))
            datos = [decodificar(v, codigos, paises if (tipo == "Países" and k == 3) else None)
                     for k, v in enumerate(fila[1:len(TIPOS_TABLA[tipo]) + 1], start=1)]
            tablas.append({"tipo": tipo, "tabla": hoja, "nombre": nombre_por_n.get(n, f"(N° {n} sin tarifa)"),
                           "datos": datos})
    return {"cabecera": cabecera[2:], "secciones": secciones[2:], "tarifas": tarifas, "tablas": tablas}


# Servicios cuyas tablas de velocidades tienen UNA sola fila (tecnología «Única»): van como
# columnas en la hoja del servicio (19/10), no en «Tablas – Velocidades».
VELOCIDADES_EN_LINEA = {
    "IF": ["Velocidades de navegación", "Velocidades con degradación"],
    "PIF": ["Velocidades de navegación"],
}
COLUMNAS_VELOCIDAD = ["Tecnología", "Descarga contratada (Mbps)", "Descarga garantizada (Mbps)",
                      "Carga contratada (Mbps)", "Carga garantizada (Mbps)"]
OBLIGATORIAS_EN_LINEA = {("IF", "Velocidades de navegación")}   # PIF y degradación: opcionales (JSON)


def velocidades_en_linea(codigo: str, d: dict) -> None:
    """Pasa las tablas de velocidades de una fila a columnas de la hoja del servicio."""
    tablas = VELOCIDADES_EN_LINEA.get(codigo)
    if not tablas:
        return
    nuevas = [f"{t}: {c}" for t in tablas for c in COLUMNAS_VELOCIDAD]
    pos = next((k for k, c in enumerate(d["cabecera"]) if c.startswith("Condiciones")), len(d["cabecera"]))
    d["cabecera"][pos:pos] = nuevas
    d["secciones"][pos:pos] = ["Características específicas"] * len(nuevas)
    por_nombre = {t["nombre"]: t for t in d["tarifas"]}
    quedan = []
    for fila in d["tablas"]:
        tarifa = por_nombre.get(fila["nombre"])
        ya = tarifa and tarifa["valores"].get(f"{fila['tabla']}: Tecnología")
        if fila["tabla"] in tablas and tarifa and not ya:
            for c, v in zip(COLUMNAS_VELOCIDAD, fila["datos"]):
                tarifa["valores"][f"{fila['tabla']}: {c}"] = v
        else:
            quedan.append(fila)                         # filas extra: siguen en la hoja de tablas
    d["tablas"] = quedan


def quitar_instalacion_si_no_aplica(codigo: str, d: dict) -> None:
    """La plantilla de SIRT trae «Tarifa instalación» en todos los servicios, pero solo la
    piden los que registran instalación (catalogo.CON_INSTALACION). En TF, móviles, etc.
    el campo sale bloqueado: se quita la columna."""
    if codigo in catalogo.CON_INSTALACION or "Tarifa instalación" not in d["cabecera"]:
        return
    k = d["cabecera"].index("Tarifa instalación")
    del d["cabecera"][k]
    del d["secciones"][k]
    for t in d["tarifas"]:
        t["valores"].pop("Tarifa instalación", None)


# Canales de TV de paga: SIRT los calcula desde un Excel subido (la plantilla de carga no
# los trae). En el Excel de registro va el NOMBRE del archivo; los archivos se guardan en
# automatizacion/canales/ (el de registro y el «_cambio» para rectificación / cambio).
CANALES = ["Canales analógicos", "Canales digitales definición estándar", "Canales digitales alta definición"]
CON_CANALES = {"TVP", "TVPR", "PTV"}
ARCHIVO_CANALES = "Canales analógicos_1.xlsx"
# Valores de muestra para obligatorios que la plantilla dejó vacíos.
RELLENOS = {"OTF": {"Tipo de línea": ["Postpago", "Prepago", "Postpago + Prepago"]}}


def agregar_canales(codigo: str, d: dict) -> None:
    if codigo not in CON_CANALES:
        return
    nuevas = [f"{c}: archivo Excel" for c in CANALES]
    pos = next((k for k, c in enumerate(d["cabecera"]) if c.startswith("Decodificadores")), len(d["cabecera"]))
    d["cabecera"][pos:pos] = nuevas
    d["secciones"][pos:pos] = ["Características específicas"] * len(nuevas)
    for t in d["tarifas"]:
        for c in nuevas:
            t["valores"][c] = ARCHIVO_CANALES


def rellenar(codigo: str, d: dict) -> None:
    for columna, opciones in RELLENOS.get(codigo, {}).items():
        for k, t in enumerate(d["tarifas"]):
            if not t["valores"].get(columna):
                t["valores"][columna] = opciones[k % len(opciones)]


# --- Listas desplegables (19/10: «que las listas de la plataforma estén en el Excel») ---
# Opciones de cada columna: del JSON del paso 2 de SIRT; si el servicio no tiene JSON
# (IF), de la plantilla de carga. Listas de varias opciones: con hasta 4 opciones se
# ofrecen todas las combinaciones («A + B»); con más, cada opción sola y se permite
# escribir combinaciones (aviso, no bloqueo).
RESPALDO_LISTAS = {
    "TF": {"Periodicidad": (["Diario", "Mensual", "Trimestral", "Semestral", "Anual"], False)},
    "IF": {"Periodicidad": (["Diario", "Mensual", "Trimestral", "Semestral", "Anual"], False),
           "Tipo de Acceso (IF)": (["Alta Nueva", "Migración"], True),
           "Tecnología (IF)": (["FTTH", "4G", "4.5G", "5G", "Satelital", "VHDL", "HFC"], False),
           "Internet Ilimitado": (["Si", "No"], False)},
}
FIJAS = {"Alcance de la tarifa": (["Residencial", "Comercial"], True)}
MULTIPLES = ("multiple", "múltiple")


def listas_json(archivo: str) -> dict:
    ruta = PROYECTO / "grabaciones" / archivo / "paso2.json"
    if not ruta.exists() or ruta.stat().st_size == 0:
        return {}
    datos = json.loads(ruta.read_text(encoding="utf-8"))["data"]
    salida = {}
    for seccion in datos.get("secciones") or []:
        for apartado in seccion.get("apartados") or []:
            opciones = [o["nombre"].strip() for o in apartado.get("lista") or [] if o.get("nombre")]
            if opciones:
                tipo = catalogo._sin_tildes((apartado.get("tipoDato") or {}).get("nombre", ""))
                salida[catalogo._sin_tildes(apartado["nombre"])] = (opciones, any(m in tipo for m in ("multiple",)))
    return salida


def opciones_de(columna: str, codigo: str, json_listas: dict):
    if columna in FIJAS:
        return FIJAS[columna]
    if columna.endswith("¿Es ilimitado?"):
        return ["Sí", "No"], False
    if codigo in VELOCIDADES_EN_LINEA and columna.endswith(": Tecnología"):
        return ["Única"], False
    if columna in RESPALDO_LISTAS.get(codigo, {}):
        return RESPALDO_LISTAS[codigo][columna]
    return json_listas.get(catalogo._sin_tildes(columna))


def combinaciones(opciones: list[str]) -> list[str]:
    from itertools import combinations
    return [" + ".join(c) for n in range(1, len(opciones) + 1) for c in combinations(opciones, n)]


TARIFAS_SIN_PLANTILLA = 10
PASO1_MUESTRA = {"Enlace Web": "https://www.prueba.com", "Fono Atención al cliente": "014567890",
                 "Email de atención al cliente": "atencion@prueba.pe"}


def desde_herramienta(servicio: dict) -> dict:
    """AC y OS (sin plantilla de muestra): 10 tarifas con los valores de la herramienta,
    variando alcance, moneda y periodicidad, con las mismas columnas del paso 1 que el resto."""
    import generar_excel as base
    modulo = importlib.import_module(f"servicios.{servicio['archivo']}")
    casos = comun.generar_casos(modulo.DIMENSIONES, modulo.CODIGO, "CM")
    paso = max(1, len(casos) // TARIFAS_SIN_PLANTILLA)
    elegidos = [casos[(k * paso + k) % len(casos)] for k in range(TARIFAS_SIN_PLANTILLA)]
    tarifas = []
    for k, caso in enumerate(elegidos, start=1):
        attrs = base.atributos(modulo, caso)
        valores = {"Alcance de la tarifa": attrs.pop("Alcance", ""), "Inicio de vigencia": "",
                   "Inicio comercialización": ""}
        if "Tarifa de instalación" in attrs:
            valores["Tarifa instalación"] = attrs.pop("Tarifa de instalación")
        valores.update(PASO1_MUESTRA)
        valores["Otros datos generales"] = attrs.pop("Otros datos generales", "")
        valores.update(attrs)
        tarifas.append({"nombre": f"TE-{modulo.CODIGO}-CM-{k:03d}", "original": "(valores de la herramienta)",
                        "valores": valores})
    cabecera = list(tarifas[0]["valores"])
    secciones = []
    for c in cabecera:
        if c in PASO_1 or c == "Tarifa instalación":
            secciones.append("Datos Generales")
        elif re.match(r"^(Condiciones|Otras condiciones)", c):
            secciones.append("Características adicionales")
        elif re.match(r"^(Restricciones|Otras restricciones)", c):
            secciones.append("Restricciones")
        else:
            secciones.append("Características específicas")
    return {"cabecera": cabecera, "secciones": secciones, "tarifas": tarifas, "tablas": []}


def generar() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    lado = Side(style="thin", color="BFBFBF")
    BORDE = Border(left=lado, right=lado, top=lado, bottom=lado)
    OBLIG = PatternFill("solid", fgColor="1F4E78")
    OPCIONAL = PatternFill("solid", fgColor="BDD7EE")
    META_F = PatternFill("solid", fgColor="375623")
    SCRIPT = PatternFill("solid", fgColor="808080")
    CELDA_SCRIPT = PatternFill("solid", fgColor="EDEDED")
    SECCION = PatternFill("solid", fgColor="DDEBF7")

    def cabeza(c, relleno, color="FFFFFF"):
        c.fill, c.border = relleno, BORDE
        c.font = Font(name=F, bold=True, size=9, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    hoja_listas = None
    rangos: dict[tuple, str] = {}

    def rango_de(opciones: list[str]) -> str:
        """Escribe la lista en la hoja oculta «Listas» (una columna por lista) y devuelve su rango."""
        nonlocal hoja_listas
        clave = tuple(opciones)
        if clave not in rangos:
            col = len(rangos) + 1
            for i, o in enumerate(opciones, start=1):
                hoja_listas.cell(i, col, o)
            letra = get_column_letter(col)
            rangos[clave] = f"Listas!${letra}$1:${letra}${len(opciones)}"
        return rangos[clave]

    def lista(hoja, col, opciones, primera=6, multiple=False):
        """Lista desplegable. Varias opciones: combinaciones si son pocas; si no, cada una
        sola y se acepta escribir «A + B» (solo aviso)."""
        valores = combinaciones(opciones) if multiple and len(opciones) <= 4 else list(opciones)
        dv = DataValidation(type="list", formula1="=" + rango_de(valores), allow_blank=True)
        if multiple and len(opciones) > 4:
            dv.errorStyle = "warning"
            dv.errorTitle = "Varias opciones"
            dv.error = "Para elegir varias, escríbelas separadas por « + » (p. ej. Whatsapp + TikTok)."
        dv.showErrorMessage = True
        hoja.add_data_validation(dv)
        letra = get_column_letter(col)
        dv.add(f"{letra}{primera}:{letra}{primera + FILAS_CON_LISTAS - 1}")

    servicios = {s["codigo"]: s for s in catalogo.servicios()}
    datos = {}
    for carpeta, codigo in CARPETAS.items():
        rutas = glob.glob(str(ORIGEN / carpeta / "*(C).xlsx"))
        if rutas:
            datos[codigo] = leer_plantilla(Path(rutas[0]), codigo)
            velocidades_en_linea(codigo, datos[codigo])
            quitar_instalacion_si_no_aplica(codigo, datos[codigo])
            agregar_canales(codigo, datos[codigo])
            rellenar(codigo, datos[codigo])
    for codigo in ORDEN:
        if codigo not in datos:
            datos[codigo] = desde_herramienta(servicios[codigo])

    wb = Workbook()
    guia = wb.active
    guia.title = "Instrucciones"
    hoja_listas = wb.create_sheet("Listas")
    hoja_listas.sheet_state = "hidden"
    resumen, todas_tablas = [], []
    listas_por_servicio: dict[str, list] = {}

    for codigo in ORDEN:
        d = datos[codigo]
        servicio = servicios[codigo]
        origen = servicio["archivo"] if (PROYECTO / "grabaciones" / servicio["archivo"] / "paso2.json").exists() \
            else catalogo.RESPALDO.get(codigo, (servicio["archivo"],))[0]
        tabla_oblig = obligatoriedad_json(origen)
        meta = ["N°", "Ejecutar", "Nombre de la tarifa", "Registro definitivo"] + \
               (["Disponibilidad comercial"] if codigo in CON_DISPONIBILIDAD else [])
        script = ["Estado", "Código SIRT", "Fecha de registro", "Detalle"]
        cab = meta + d["cabecera"] + script
        secciones = [""] * len(meta) + d["secciones"] + ["Resultado"] * len(script)

        hoja = wb.create_sheet(codigo)
        hoja["A1"] = f"{codigo} · {servicio['nombre']}"
        hoja["A1"].font = Font(name=F, bold=True, size=14)
        hoja["A2"] = ("Una fila por tarifa (agrega las que necesites). Fechas vacías = día del registro. "
                      "Las tablas de esta tarifa van en las hojas «Tablas – …», unidas por el nombre.")
        hoja["A2"].font = Font(name=F, italic=True, size=9, color="595959")
        for j, (texto, relleno, color) in enumerate([("Obligatorio", OBLIG, "FFFFFF"), ("Opcional", OPCIONAL, "1F1F1F"),
                                                     ("Lo llena el script", SCRIPT, "FFFFFF")], start=4):
            cabeza(hoja.cell(3, j, texto), relleno, color)
        # Fila 4: sección (Datos generales / Características específicas / …)
        for j, s in enumerate(secciones, start=1):
            c = hoja.cell(4, j, s if (j == 1 or s != secciones[j - 2]) else None)
            c.fill, c.border = SECCION, BORDE
            c.font = Font(name=F, bold=True, size=8, color="1F4E78")
        for j, titulo in enumerate(cab, start=1):
            if j <= len(meta):
                cabeza(hoja.cell(5, j, titulo), META_F)
            elif j > len(cab) - len(script):
                cabeza(hoja.cell(5, j, titulo), SCRIPT)
            elif es_obligatorio(titulo, tabla_oblig, codigo):
                cabeza(hoja.cell(5, j, titulo.strip()), OBLIG)
            else:
                cabeza(hoja.cell(5, j, titulo.strip()), OPCIONAL, "1F1F1F")
        hoja.row_dimensions[5].height = 54

        for i, t in enumerate(d["tarifas"], start=6):
            v = dict(t["valores"])
            for fecha in ("Inicio de vigencia", "Inicio comercialización"):
                if fecha in v:
                    v[fecha] = ""                            # vacío = día del registro
            fila = [i - 5, "Sí", t["nombre"], "Sí"] + ([DISPONIBILIDAD_POR_DEFECTO] if codigo in CON_DISPONIBILIDAD else []) \
                + [v.get(c, "") for c in d["cabecera"]] + [""] * len(script)
            for j, valor in enumerate(fila, start=1):
                c = hoja.cell(i, j, valor)
                c.font = Font(name=F, size=9)
                c.border = BORDE
                c.alignment = Alignment(vertical="top", wrap_text=j > 3,
                                        horizontal="center" if j in (1, 2, 4) else None)
                if j > len(cab) - len(script):
                    c.fill = CELDA_SCRIPT

        lista(hoja, 2, ["Sí", "No"])
        lista(hoja, 4, ["Sí", "No"])
        json_listas = listas_json(origen)
        con_lista = []
        for j, titulo in enumerate(cab, start=1):
            if j <= len(meta) or j > len(cab) - len(script):
                continue
            encontrada = opciones_de(titulo, codigo, json_listas)
            if encontrada:
                lista(hoja, j, encontrada[0], multiple=encontrada[1])
                con_lista.append(titulo)
        listas_por_servicio[codigo] = con_lista
        anchos = [5, 9, 24, 11] + ([14] if codigo in CON_DISPONIBILIDAD else []) + \
                 [max(12, min(40, len(c) * 0.9 + 4)) for c in d["cabecera"]] + [12, 18, 16, 40]
        for j, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(j)].width = ancho
        hoja.freeze_panes = "D6"
        hoja.auto_filter.ref = f"A5:{get_column_letter(len(cab))}{5 + len(d['tarifas'])}"
        for t in d["tablas"]:
            todas_tablas.append((codigo, t))
        oblig = [es_obligatorio(c, tabla_oblig, codigo) for c in d["cabecera"]]
        resumen.append((codigo, servicio["nombre"], len(d["tarifas"]), len(d["cabecera"]),
                        sum(oblig), len(oblig) - sum(oblig), sum(1 for t in d["tablas"]),
                        "plantilla SIRT (C)" if codigo in CARPETAS.values() else "valores de la herramienta"))

    # Hojas de tablas: una por tipo.
    TABLAS_POR_TIPO = {
        "Países": ["Tarifa por minuto a fijo LDI", "Tarifa por minuto a móvil LDI", "Tarifa por minuto LDI",
                   "Tarifa minuto roaming", "Tarifa SMS roaming", "Tarifa MB roaming", "Minutos LDI países",
                   "Datos Roaming del plan países", "Datos Roaming adicional países", "Minutos Roaming países",
                   "SMS Roaming países"],
        "Velocidades": ["Tecnologías y velocidades"],
        "Bolsas": ["Bolsa de datos para apps", "Datos para Apps 1", "Datos para Apps 2", "Datos para Apps 3",
                   "Datos para Apps 4"],
    }
    from openpyxl import load_workbook
    PAISES = []
    muestra = glob.glob(str(ORIGEN / "SM POS" / "*(C).xlsx"))
    if muestra:
        PAISES = sorted(maestras(load_workbook(muestra[0], data_only=True))[1].values())
    for tipo, columnas in TIPOS_TABLA.items():
        hoja = wb.create_sheet(f"Tablas – {tipo}")
        hoja["A1"] = f"Tablas de {tipo.lower()}"
        hoja["A1"].font = Font(name=F, bold=True, size=14)
        hoja["A2"] = ("Una fila por fila de tabla. «Nombre de la tarifa» debe coincidir con el de la hoja de su servicio. "
                      "Varias opciones (países, plataformas) separadas por « + ».")
        hoja["A2"].font = Font(name=F, italic=True, size=9, color="595959")
        cab = ["Servicio", "Nombre de la tarifa", "Tabla"] + columnas
        for j, titulo in enumerate(cab, start=1):
            cabeza(hoja.cell(5, j, titulo), META_F if j <= 3 else OBLIG)
        hoja.row_dimensions[5].height = 40
        filas = [(c, t) for c, t in todas_tablas if t["tipo"] == tipo]
        for i, (codigo, t) in enumerate(filas, start=6):
            for j, valor in enumerate([codigo, t["nombre"], t["tabla"]] + t["datos"], start=1):
                c = hoja.cell(i, j, valor)
                c.font, c.border = Font(name=F, size=9), BORDE
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for j, ancho in enumerate([9, 22, 30] + [16 if k else 18 for k in range(len(columnas))], start=1):
            hoja.column_dimensions[get_column_letter(j)].width = ancho
        if tipo == "Países":
            hoja.column_dimensions[get_column_letter(len(cab))].width = 40
        hoja.freeze_panes = "D6"
        hoja.auto_filter.ref = f"A5:{get_column_letter(len(cab))}{max(6, 5 + len(filas))}"
        lista(hoja, 1, ["SMP", "SMC", "SMO", "SMQ"] if tipo != "Países" else ["TF", "TFP", "PTF", "SMP", "SMC", "SMO", "SMQ"])
        lista(hoja, 3, TABLAS_POR_TIPO[tipo])
        if tipo == "Países":
            lista(hoja, 4, ["Bloque"])
            lista(hoja, 6, PAISES or ["ALEMANIA"], multiple=True)
        elif tipo == "Velocidades":
            lista(hoja, 4, ["2G", "3G", "4G", "4.5G", "5G"])
        else:
            lista(hoja, 4, ["Bolsa"])
            lista(hoja, 6, ["Otros", "Whatsapp", "Messenger", "Facebook", "TikTok", "YouTube", "Instagram"], multiple=True)

    textos = [
        ("Registro de tarifas establecidas", True, 14),
        ("Datos de muestra: plantillas de carga de SIRT (C) de 13 servicios; AC y OS (10 c/u) con valores de la herramienta. Ubicación: automatizacion/.", False, 9),
        ("", False, 10),
        ("Cómo se usa", True, 11),
        ("1. Abre la hoja del servicio (código en la pestaña) y escribe una fila por tarifa: 1, 2, 3… las que quieras.", False, 10),
        ("   Solo se registran los servicios cuya hoja está en el archivo: para registrar uno solo, deja solo su hoja", False, 10),
        ("   (y, si tiene tablas, la hoja «Tablas – …» que corresponda). No borres las hojas «Instrucciones» ni «Listas».", False, 10),
        ("2. «Ejecutar» = Sí: se procesa; = No: se deja en el Excel sin tocar (estado «Omitida»).", False, 10),
        ("3. El nombre de la tarifa debe ser único: es lo que usa el script para no registrar dos veces.", False, 10),
        ("4. Si la tarifa tiene tablas (países, velocidades, bolsas), agrega sus filas en la hoja «Tablas – …» con el mismo nombre.", False, 10),
        ("5. Guarda, cierra el Excel y corre el script. Las columnas grises mostrarán el resultado de cada fila.", False, 10),
        ("", False, 10),
        ("Colores", True, 11),
        ("Verde: datos de la fila. Azul oscuro: obligatorio en SIRT. Azul claro: opcional. Gris: lo llena el script.", False, 10),
        ("", False, 10),
        ("Valores", True, 11),
        ("Varias opciones: separadas por « + » (Residencial + Comercial; Alta nueva + Migración).", False, 10),
        ("Minutos/SMS: «¿Es ilimitado?» Sí/No; si es No, completar cantidad y tarifa adicional (por eso esas dos van en azul claro).", False, 10),
        ("Canales (TV de paga): nombre del archivo Excel guardado en automatizacion/canales/.", False, 10),
        ("Listas desplegables: las columnas que en SIRT son listas tienen sus opciones (flecha en la celda).", False, 10),
        ("  Si admiten varias opciones y son pocas, la lista trae todas las combinaciones (A + B).", False, 10),
        ("  Si son muchas (plataformas, países), elige una o escribe varias separadas por « + ».", False, 10),
        (f"Disponibilidad comercial (solo fijos, internet y TV): por defecto {DISPONIBILIDAD_POR_DEFECTO}.", False, 10),
        ("Inicio de vigencia e inicio de comercialización son obligatorios en SIRT; si se dejan vacíos, el script usa el día del registro.", False, 10),
        ("", False, 10),
        ("Servicios", True, 11),
    ]
    for i, (t, negrita, tam) in enumerate(textos, start=1):
        guia.cell(i, 1, t).font = Font(name=F, bold=negrita, size=tam)
    fila = len(textos) + 1
    titulos = ["Hoja", "Servicio", "Tarifas", "Atributos", "Obligatorios", "Opcionales", "Filas de tablas", "Origen"]
    for j, titulo in enumerate(titulos, start=1):
        c = guia.cell(fila, j, titulo)
        c.font, c.fill = Font(name=F, bold=True, color="FFFFFF"), OBLIG
    for r in resumen:
        fila += 1
        for j, v in enumerate(r, start=1):
            guia.cell(fila, j, v).font = Font(name=F, size=10)
    for col, ancho in zip("ABCDEFGH", [10, 56, 9, 10, 12, 11, 14, 24]):
        guia.column_dimensions[col].width = ancho

    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Generado: {SALIDA}")
    for r in resumen:
        print(f"  {r[0]:<5} {r[2]:>3} tarifas · {r[3]:>3} atributos ({r[4]} oblig., {r[5]} opc.) · {r[6]:>3} filas de tablas · {r[7]}")


if __name__ == "__main__":
    generar()

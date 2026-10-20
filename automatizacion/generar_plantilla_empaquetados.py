"""
Genera automatizacion/empaquetados.xlsx (BORRADOR, sin conectar con los scripts): el
Excel para registrar tarifas EMPAQUETADAS a partir de tarifas monoproducto ya
registradas (con registro definitivo).

Una fila = un paquete. Las tarifas base van en columnas «Tarifa 1 … Tarifa 6», cada
una con su «Valor en el paquete» (vacío = valor de la tarifa base − 1, como hoy).

Reglas de SIRT que el script validará antes de registrar (vistas en pruebas):
- 2 o más tarifas, todas con registro definitivo y del MISMO alcance que el paquete.
- No puede haber dos tarifas del mismo servicio (cada una es una pestaña «Servicio N»).
- Valor del plan = suma de los valores: debe ser MENOR que la suma de las tarifas base
  y MAYOR O IGUAL que el valor de la tarifa base más cara.

Uso (desde la carpeta del proyecto):
    python automatizacion/generar_plantilla_empaquetados.py
"""

import sys
from pathlib import Path

PROYECTO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROYECTO))

SALIDA = PROYECTO / "automatizacion" / "empaquetados.xlsx"
MAX_TARIFAS = 6
FILAS_CON_LISTAS = 1000

DATOS = [  # (columna, obligatorio)
    ("Alcance de la tarifa", True), ("Inicio de vigencia", True), ("Inicio comercialización", True),
    ("Tarifa instalación", False), ("Enlace Web", True), ("Fono Atención al cliente", True),
    ("Email de atención al cliente", True), ("Otros datos generales", False),
]
CONTACTO = {"Enlace Web": "https://www.prueba.com", "Fono Atención al cliente": "014567890",
            "Email de atención al cliente": "atencion@prueba.pe"}

# Muestras: tarifas del Excel de registro (automatizacion/registro_tarifas_establecidas.xlsx).
# (alcance, [(tarifa base, valor en el paquete o "" = base − 1)], otros datos)
MUESTRAS = [
    ("Residencial", [("TE-TFP-CM-001", ""), ("TE-SMP-CM-001", "")], "Dúo fijo + móvil"),
    ("Residencial", [("TE-TF-CM-001", ""), ("TE-IF-CM-003", ""), ("TE-TVP-CM-001", "")], "Trío fijo + internet + TV"),
    ("Residencial", [("TE-SMC-CM-001", "18.9"), ("TE-OS-CM-001", "45")], "Valores escritos a mano"),
    ("Comercial", [("TE-TFP-CM-002", ""), ("TE-TVP-CM-002", ""), ("TE-SMP-CM-004", ""), ("TE-AC-CM-004", "")],
     "Cuatro servicios, alcance comercial"),
    ("Residencial + Comercial", [("TE-TF-CM-003", ""), ("TE-TVP-CM-003", "")], "Alcance doble"),
]


def generar() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    lado = Side(style="thin", color="BFBFBF")
    BORDE = Border(left=lado, right=lado, top=lado, bottom=lado)
    OBLIG = PatternFill("solid", fgColor="1F4E78")
    OPCIONAL = PatternFill("solid", fgColor="BDD7EE")
    META_F = PatternFill("solid", fgColor="375623")
    SCRIPT = PatternFill("solid", fgColor="808080")
    CELDA_SCRIPT = PatternFill("solid", fgColor="EDEDED")
    CALCULO = PatternFill("solid", fgColor="E2EFDA")
    SECCION = PatternFill("solid", fgColor="DDEBF7")

    wb = Workbook()
    guia = wb.active
    guia.title = "Instrucciones"
    hoja = wb.create_sheet("Empaquetados")
    listas = wb.create_sheet("Listas")
    listas.sheet_state = "hidden"

    meta = [("N°", None), ("Ejecutar", None), ("Nombre del paquete", None), ("Registro definitivo", None),
            ("Disponibilidad comercial", None)]
    tarifas = []
    for k in range(1, MAX_TARIFAS + 1):
        tarifas += [(f"Tarifa {k}", k <= 2), (f"Valor {k} en el paquete", False)]
    calculo = [("Cantidad de tarifas", None)]
    script = [("Estado", None), ("Código SIRT", None), ("Fecha de registro", None), ("Detalle", None)]
    columnas = meta + DATOS + tarifas + calculo + script
    secciones = [""] * len(meta) + ["Datos Generales"] * len(DATOS) + \
        ["Tarifas del paquete (nombre de la tarifa base y su valor)"] * len(tarifas) + [""] + ["Resultado"] * len(script)

    hoja["A1"] = "Tarifas empaquetadas"
    hoja["A1"].font = Font(name=F, bold=True, size=14)
    hoja["A2"] = ("Una fila por paquete. Tarifas base: nombre exacto de tarifas ya registradas en definitivo. "
                  "Valor vacío = valor de la tarifa base − 1.")
    hoja["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    for j, (texto, relleno, color) in enumerate([("Obligatorio", OBLIG, "FFFFFF"), ("Opcional", OPCIONAL, "1F1F1F"),
                                                 ("Calculado", CALCULO, "1F1F1F"),
                                                 ("Lo llena el script", SCRIPT, "FFFFFF")], start=4):
        c = hoja.cell(3, j, texto)
        c.fill, c.border = relleno, BORDE
        c.font = Font(name=F, bold=True, size=9, color=color)
        c.alignment = Alignment(horizontal="center")

    for j, s in enumerate(secciones, start=1):
        c = hoja.cell(4, j, s if (s and (j == 1 or s != secciones[j - 2])) else None)
        c.fill, c.border = SECCION, BORDE
        c.font = Font(name=F, bold=True, size=8, color="1F4E78")

    for j, (titulo, oblig) in enumerate(columnas, start=1):
        c = hoja.cell(5, j, titulo)
        if j <= len(meta):
            relleno, color = META_F, "FFFFFF"
        elif (titulo, oblig) in script:
            relleno, color = SCRIPT, "FFFFFF"
        elif (titulo, oblig) in calculo:
            relleno, color = CALCULO, "1F1F1F"
        elif oblig:
            relleno, color = OBLIG, "FFFFFF"
        else:
            relleno, color = OPCIONAL, "1F1F1F"
        c.fill, c.border = relleno, BORDE
        c.font = Font(name=F, bold=True, size=9, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    hoja.row_dimensions[5].height = 45

    primera_tarifa = len(meta) + len(DATOS) + 1
    col_cantidad = primera_tarifa + len(tarifas)
    letras_tarifa = [get_column_letter(primera_tarifa + 2 * k) for k in range(MAX_TARIFAS)]
    for i, (alcance, componentes, otros) in enumerate(MUESTRAS, start=6):
        n = i - 5
        datos = {"Alcance de la tarifa": alcance, "Inicio de vigencia": "", "Inicio comercialización": "",
                 "Tarifa instalación": "", **CONTACTO, "Otros datos generales": otros}
        fila = [n, "Sí", f"TE-EMP-CM-{n:03d}", "Sí", "Lima"] + [datos[c] for c, _ in DATOS]
        for k in range(MAX_TARIFAS):
            fila += list(componentes[k]) if k < len(componentes) else ["", ""]
        fila += [None] + [""] * len(script)
        for j, v in enumerate(fila, start=1):
            c = hoja.cell(i, j, v)
            c.font = Font(name=F, size=9)
            c.border = BORDE
            c.alignment = Alignment(vertical="top", wrap_text=j > 3,
                                    horizontal="center" if j in (1, 2, 4) else None)
            if j > col_cantidad:
                c.fill = CELDA_SCRIPT
    # «Cantidad de tarifas»: fórmula para las filas con datos y las siguientes (vacías).
    for i in range(6, 6 + FILAS_CON_LISTAS):
        celdas = ",".join(f"{l}{i}" for l in letras_tarifa)
        c = hoja.cell(i, col_cantidad, f'=IF(C{i}="","",COUNTA({celdas}))')
        c.fill, c.border = CALCULO, BORDE
        c.font = Font(name=F, size=9)
        c.alignment = Alignment(horizontal="center", vertical="top")

    # Listas desplegables (opciones en la hoja oculta «Listas»).
    opciones = {"A": ["Sí", "No"], "B": ["Residencial", "Comercial", "Residencial + Comercial"]}
    for letra, valores in opciones.items():
        for k, v in enumerate(valores, start=1):
            listas[f"{letra}{k}"] = v
    for col, rango in ((2, "A"), (4, "A"), (len(meta) + 1, "B")):
        dv = DataValidation(type="list", formula1=f"=Listas!${rango}$1:${rango}${len(opciones[rango])}",
                            allow_blank=True)
        hoja.add_data_validation(dv)
        letra = get_column_letter(col)
        dv.add(f"{letra}6:{letra}{5 + FILAS_CON_LISTAS}")

    anchos = [5, 9, 20, 11, 14, 22, 12, 12, 11, 22, 13, 22, 26] + [18, 11] * MAX_TARIFAS + [10, 12, 18, 16, 40]
    for j, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(j)].width = ancho
    hoja.freeze_panes = "D6"
    hoja.auto_filter.ref = f"A5:{get_column_letter(len(columnas))}{5 + len(MUESTRAS)}"

    textos = [
        ("Registro de tarifas empaquetadas (BORRADOR)", True, 14),
        ("Ubicación: automatizacion/empaquetados.xlsx. Las tarifas base salen del Excel de registro "
         "(registro_tarifas_establecidas.xlsx) y deben tener registro definitivo.", False, 9),
        ("", False, 10),
        ("Cómo se usa", True, 11),
        ("1. Una fila por paquete (agrega las que necesites). «Ejecutar» = No: la fila se deja sin tocar.", False, 10),
        ("2. Escribe el nombre exacto de cada tarifa base en «Tarifa 1», «Tarifa 2»… (mínimo 2, máximo 6).", False, 10),
        ("3. «Valor N en el paquete»: lo que cobra ese servicio dentro del paquete. Vacío = valor de la tarifa base − 1.", False, 10),
        ("4. Guarda, cierra el Excel y corre el script. Las columnas grises mostrarán el resultado de cada paquete.", False, 10),
        ("", False, 10),
        ("Reglas de SIRT (el script las revisa antes de registrar)", True, 11),
        ("• Todas las tarifas base deben tener registro definitivo y el MISMO alcance que el paquete.", False, 10),
        ("• No puede haber dos tarifas del mismo servicio en un paquete.", False, 10),
        ("• Suma de los valores < suma de las tarifas base, y ≥ valor de la tarifa base más cara.", False, 10),
        ("  Ejemplo: bases 60 y 29.9 → valores 59 y 28.9 → suma 87.9 (menor que 89.9 y mayor que 60): válido.", False, 10),
        ("• Disponibilidad comercial (por defecto Lima): solo se usa si alguna tarifa base es de fijos, internet o TV.", False, 10),
        ("• Tarifa instalación: solo si SIRT la habilita para el paquete; vacía = 10.", False, 10),
        ("• Fechas de inicio vacías: se usa el día del registro.", False, 10),
        ("", False, 10),
        ("Colores", True, 11),
        ("Verde oscuro: datos de la fila. Azul oscuro: obligatorio. Azul claro: opcional. Verde claro: calculado. "
         "Gris: lo llena el script.", False, 10),
    ]
    for i, (t, negrita, tam) in enumerate(textos, start=1):
        guia.cell(i, 1, t).font = Font(name=F, bold=negrita, size=tam)
    guia.column_dimensions["A"].width = 120

    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Generado: {SALIDA} ({len(MUESTRAS)} paquetes de muestra)")


if __name__ == "__main__":
    generar()

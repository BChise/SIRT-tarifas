"""
Genera automatizacion/modificaciones.xlsx: el Excel para RECTIFICACIÓN y CAMBIO DE TARIFA de
monoproductos, empaquetados y licitaciones (lo lee modificaciones_excel.py).

Se indica QUÉ hacer (reglas), no el valor nuevo de cada campo: el script toma lo que
la tarifa tiene en SIRT y lo modifica según las opciones de la fila (lo mismo que hoy
hace acciones.py rectificar / cambio, ya probado).

Uso (desde la carpeta del proyecto):
    python automatizacion/generar_plantilla_modificaciones.py
"""

from pathlib import Path

PROYECTO = Path(__file__).resolve().parent.parent
SALIDA = PROYECTO / "automatizacion" / "modificaciones.xlsx"
FILAS_CON_LISTAS = 1000

ACCIONES = ["Rectificación", "Cambio"]
CAMPOS = ["Por defecto", "Todos los campos", "Solo la lista", "Detalle en hoja Campos"]
SENTIDOS = ["Beneficioso", "Perjudicial"]
LISTAS = {"Sí/No": ["Sí", "No"], "Acción": ACCIONES, "Campos": CAMPOS, "Sentido": SENTIDOS}

# (título, tipo, lista, ancho, sección)
COLUMNAS = [
    ("N°", "meta", None, 5, ""), ("Ejecutar", "meta", "Sí/No", 9, ""),
    ("Nombre de la tarifa", "meta", None, 22, ""),
    ("Acción", "obligatorio", "Acción", 14, "Qué hacer"),
    ("Delta (vacío: +1 rectificación, −1 cambio)", "opcional", None, 13, "Qué hacer"),
    ("Cambio: sentido (Beneficioso / Perjudicial)", "opcional", "Sentido", 14, "Qué hacer"),
    ("Campos a modificar", "obligatorio", "Campos", 16, "Qué hacer"),
    ("Lista de campos (separados por ;)", "opcional", None, 30, "Qué hacer"),
    ("Agregar una fila a cada tabla", "opcional", "Sí/No", 11, "Tablas y condicionales (paso 2)"),
    ("Pasar «ilimitado» a limitado", "opcional", "Sí/No", 11, "Tablas y condicionales (paso 2)"),
    ("Cambiar canales (TV de paga)", "opcional", "Sí/No", 11, "Tablas y condicionales (paso 2)"),
    ("Cambio: fecha de ejecución", "opcional", None, 13, "Solo cambio"),
    ("Texto del sello (vacío: «Rectificación» / «Cambio»)", "opcional", None, 18, "Textos"),
    ("Estado", "script", None, 13, "Resultado"), ("Código SIRT", "script", None, 18, "Resultado"),
    ("Fecha de ejecución", "script", None, 17, "Resultado"), ("Detalle", "script", None, 60, "Resultado"),
]

MUESTRAS = [  # (tarifa, acción, delta, sentido, campos, lista, agregar, ilimitado, canales, fecha)
    ("TE-IF-CM-008", "Rectificación", "", "", "Por defecto", "", "No", "No", "", ""),
    ("TE-SMP-CM-009", "Cambio", "1", "Beneficioso", "Todos los campos", "", "Sí", "No", "", ""),
    ("TE-SMC-CM-009", "Cambio", "1", "Perjudicial", "Todos los campos", "", "No", "Sí", "", ""),
    ("TE-TF-CM-008", "Rectificación", "", "", "Solo la lista", "Otros datos generales; Valor del servicio", "No", "No", "", ""),
    ("TE-SMQ-CM-009", "Cambio", "", "", "Detalle en hoja Campos", "", "No", "No", "", ""),
]
DETALLE_CAMPOS = [  # (tarifa, campo o ruta con «›», valor nuevo, delta)
    ("TE-SMQ-CM-009", "Datos tope descarga", "", "5"),
    ("TE-SMQ-CM-009", "Datos para tethering", "3", ""),
    ("TE-SMQ-CM-009", "Valor del servicio", "", "-2"),
    ("TE-SMQ-CM-009", "Minutos móvil on net › Cantidad de minutos", "300", ""),
    ("TE-SMQ-CM-009", "Tecnologías y velocidades › 4G › Descarga máx. contratada", "", "5"),
    ("TE-SMQ-CM-009", "Tecnologías y velocidades › 5G", "", "-1"),
    ("TE-SMQ-CM-009", "Tecnologías y velocidades › 3G › Tecnología", "4.5G", ""),
]

DESCRIPCION = [
    ("Qué hace cada opción", True, 14),
    ("", False, 10),
    ("ACCIÓN", True, 12),
    ("• Rectificación: corrige la tarifa vigente; se guarda de inmediato.", False, 10),
    ("• Cambio: programa una nueva versión de la tarifa para una fecha de ejecución. Solo puede haber UN cambio", False, 10),
    ("  pendiente por tarifa: si ya hay uno programado, la fila queda «No aplica» hasta que se ejecute.", False, 10),
    ("", False, 10),
    ("CAMBIO: SENTIDO (Beneficioso / Perjudicial)", True, 12),
    ("Solo en «Cambio». Separa los campos por su nombre en PRECIOS (valor del servicio, tarifa por minuto / SMS / MB", False, 10),
    ("adicional, instalación, periodo de contrato) y BENEFICIOS (minutos, SMS, datos, GB, velocidades, duración,", False, 10),
    ("decodificadores, repetidores…). Usa el valor del Delta sin signo (vacío = 1):", False, 10),
    ("• Beneficioso para el usuario: los precios BAJAN y los beneficios SUBEN (sugerido: «Agregar una fila» = Sí).", False, 10),
    ("• Perjudicial: los precios SUBEN y los beneficios BAJAN (sugerido: «Pasar ilimitado a limitado» = Sí).", False, 10),
    ("• Vacío: el Delta se aplica igual a todos los números (como antes). En rectificación se ignora.", False, 10),
    ("", False, 10),
    ("DELTA (cuánto suben o bajan los números)", True, 12),
    ("Vacío: +1 en rectificación y −1 en cambio. Se puede poner otro número (p. ej. 2, −1, 0.5).", False, 10),
    ("Nunca baja de 0, los porcentajes no pasan de 100 y los «topes» se cambian en el orden que exige SIRT.", False, 10),
    ("Empaquetados que suben: el último servicio no sube (regla: suma de servicios < suma de las tarifas base).", False, 10),
    ("", False, 10),
    ("CAMPOS A MODIFICAR", True, 12),
    ("• Por defecto: textos (con el sello «| Rectificación/Cambio fecha hora»), valor del servicio, tarifa de", False, 10),
    ("  instalación y listas como «Tipo de Contratación».", False, 10),
    ("• Todos los campos: además TODOS los números del paso 2 (datos, velocidades, minutos, tarifas…), «Apps sin", False, 10),
    ("  consumo de datos» y los números de cada fila de las tablas (lápiz → Guardar), sin tocar tecnología, países", False, 10),
    ("  ni plataformas (SIRT no permite repetirlos).", False, 10),
    ("• Solo la lista: solo los campos escritos en «Lista de campos», con su nombre como aparece en SIRT.", False, 10),
    ("• Detalle en hoja Campos: cada campo de la tarifa con su «Valor nuevo» o su propio «Delta» (hoja «Campos»;", False, 10),
    ("  una fila por campo, nombre como aparece en SIRT). SOLO se modifican esos campos (y los textos).", False, 10),
    ("  En «Campo» se puede escribir una ruta con « › » para ser más fino:", False, 10),
    ("    Datos tope descarga                                   → ese campo", False, 10),
    ("    Minutos móvil on net › Cantidad de minutos            → la cantidad de ese bloque (no su tarifa)", False, 10),
    ("    Tecnologías y velocidades                             → todos los números de todas las filas de la tabla", False, 10),
    ("    Tecnologías y velocidades › 4G                        → todos los números de la fila 4G", False, 10),
    ("    Tecnologías y velocidades › 4G › Descarga máx. contratada → solo esa velocidad de la fila 4G", False, 10),
    ("    Tecnologías y velocidades › 3G › Tecnología  (Valor nuevo 4.5G) → cambia la tecnología de esa fila", False, 10),
    ("    Tarifa minuto roaming › ALEMANIA › Tarifa              → la tarifa de la fila de Alemania", False, 10),
    ("  La «fila» es el texto de una celda de la tabla (tecnología, país, tipo…). SIRT no deja repetir tecnología ni", False, 10),
    ("  país en una tabla: el valor nuevo debe ser uno que la tabla aún no tenga.", False, 10),
    ("", False, 10),
    ("TABLAS Y CONDICIONALES (paso 2)", True, 12),
    ("• Agregar una fila a cada tabla: una fila nueva con un país / tecnología / plataforma que la tabla aún no tenga.", False, 10),
    ("• Pasar «ilimitado» a limitado: en minutos y SMS marcados como ilimitados, desmarca la casilla y llena", False, 10),
    ("  cantidad y tarifa adicional (10 y 15).", False, 10),
    ("• Cambiar canales: sube el otro Excel de canales de automatizacion/canales/ (3 ↔ 5 canales). Vacío = Sí.", False, 10),
    ("", False, 10),
    ("CAMBIO: FECHA DE EJECUCIÓN", True, 12),
    ("Vacío: hoy + 3 días (SIRT pide al menos 2 si el cambio perjudica al usuario; hoy siempre está bloqueado).", False, 10),
    ("Si el día está bloqueado en el calendario, usa el siguiente permitido (hasta 7 días).", False, 10),
    ("", False, 10),
    ("POR TIPO DE TARIFA", True, 12),
    ("• Monoproductos y empaquetados: todo lo anterior (en empaquetados, en cada pestaña de servicio).", False, 10),
    ("• Licitación y Contratos (solo rectificación): además mueve inicio y fin de vigencia (+1 día, con el", False, 10),
    ("  calendario), cambia Ámbito y Moneda a la otra opción y suma el delta al Valor de contrato.", False, 10),
    ("• Promociones: la rectificación / cambio de promociones aún no está grabada; esas filas quedan en «Error de datos».", False, 10),
    ("", False, 10),
    ("REGLAS GENERALES", True, 12),
    ("• La tarifa debe estar EN COMERCIALIZACIÓN y VIGENTE (licitaciones: vigente). Si no, «No aplica».", False, 10),
    ("• Antes de guardar se comprueba que todos los cambios aparezcan en el resumen de SIRT.", False, 10),
    ("• Solo se procesan las filas con «Ejecutar» = Sí. El Excel debe estar cerrado al correr el comando.", False, 10),
]


def generar() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    lado = Side(style="thin", color="BFBFBF")
    BORDE = Border(left=lado, right=lado, top=lado, bottom=lado)
    RELLENO = {"meta": ("375623", "FFFFFF"), "obligatorio": ("1F4E78", "FFFFFF"), "opcional": ("BDD7EE", "1F1F1F"),
               "script": ("808080", "FFFFFF")}
    CELDA_SCRIPT = PatternFill("solid", fgColor="EDEDED")
    SECCION = PatternFill("solid", fgColor="DDEBF7")

    def cabeza(c, tipo):
        fondo, color = RELLENO[tipo]
        c.fill, c.border = PatternFill("solid", fgColor=fondo), BORDE
        c.font = Font(name=F, bold=True, size=9, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Modificaciones"
    guia = wb.create_sheet("Qué hace cada opción")
    listas = wb.create_sheet("Listas")
    listas.sheet_state = "hidden"
    rangos = {}
    for k, (nombre, valores) in enumerate(LISTAS.items(), start=1):
        letra = get_column_letter(k)
        for f, v in enumerate(valores, start=1):
            listas[f"{letra}{f}"] = v
        rangos[nombre] = f"=Listas!${letra}$1:${letra}${len(valores)}"

    hoja["A1"] = "Rectificación y cambio de tarifa"
    hoja["A1"].font = Font(name=F, bold=True, size=14)
    hoja["A2"] = ("Una fila por acción. Se indica qué hacer, no el valor nuevo: el script toma lo que la tarifa tiene "
                  "en SIRT y lo modifica. Detalle en la hoja «Qué hace cada opción».")
    hoja["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    # Leyenda de colores: cada etiqueta sobre una columna de ese tipo (no en columnas fijas).
    tipos = [c[1] for c in COLUMNAS]
    for texto, clave in (("Obligatorio", "obligatorio"), ("Opcional", "opcional"), ("Lo llena el script", "script")):
        cabeza(hoja.cell(3, tipos.index(clave) + 1, texto), clave)
    secciones = [c[4] for c in COLUMNAS]
    for j, s in enumerate(secciones, start=1):
        c = hoja.cell(4, j, s if (s and (j == 1 or s != secciones[j - 2])) else None)
        c.fill, c.border = SECCION, BORDE
        c.font = Font(name=F, bold=True, size=8, color="1F4E78")
    for j, (titulo, tipo, nombre_lista, ancho, _) in enumerate(COLUMNAS, start=1):
        cabeza(hoja.cell(5, j, titulo), tipo)
        hoja.column_dimensions[get_column_letter(j)].width = ancho
        if nombre_lista:
            dv = DataValidation(type="list", formula1=rangos[nombre_lista], allow_blank=True)
            dv.showErrorMessage, dv.error = True, "Elige una opción de la lista."
            hoja.add_data_validation(dv)
            letra = get_column_letter(j)
            dv.add(f"{letra}6:{letra}{5 + FILAS_CON_LISTAS}")
    hoja.row_dimensions[5].height = 58

    col = {t: j for j, (t, *_r) in enumerate(COLUMNAS, start=1)}
    letra_de = lambda titulo: get_column_letter(col[titulo])
    for i, muestra in enumerate(MUESTRAS, start=6):
        fila = [i - 5, "Sí", *muestra] + [""] * (len(COLUMNAS) - 2 - len(muestra))
        for j, v in enumerate(fila, start=1):
            c = hoja.cell(i, j, v)
            c.font, c.border = Font(name=F, size=9), BORDE
            c.alignment = Alignment(vertical="top", wrap_text=True,
                                    horizontal="center" if j in (1, 2, 5, 6, 9, 10, 11) else None)
            if COLUMNAS[j - 1][1] == "script":
                c.fill = CELDA_SCRIPT
    for i in range(6, 6 + FILAS_CON_LISTAS):
        hoja.cell(i, col["Cambio: fecha de ejecución"]).number_format = "DD/MM/YYYY"
    delta = DataValidation(type="decimal", operator="between", formula1="-1000", formula2="1000", allow_blank=True)
    delta.showErrorMessage, delta.error = True, "El delta es un número (p. ej. 1, -1, 0.5)."
    hoja.add_data_validation(delta)
    delta.add(f"{letra_de('Delta (vacío: +1 rectificación, −1 cambio)')}6:{letra_de('Delta (vacío: +1 rectificación, −1 cambio)')}{5 + FILAS_CON_LISTAS}")
    fecha = DataValidation(type="date", operator="greaterThan", formula1="TODAY()", allow_blank=True)
    fecha.showErrorMessage, fecha.error = True, "La fecha de ejecución debe ser posterior a hoy (vacío = hoy + 3)."
    hoja.add_data_validation(fecha)
    fecha.add(f"{letra_de('Cambio: fecha de ejecución')}6:{letra_de('Cambio: fecha de ejecución')}{5 + FILAS_CON_LISTAS}")
    hoja.freeze_panes = "D6"
    hoja.auto_filter.ref = f"A5:{get_column_letter(len(COLUMNAS))}{5 + len(MUESTRAS)}"

    # --- Hoja Campos (modo «Detalle en hoja Campos») ----------------------------------
    hc = wb.create_sheet("Campos", 1)
    hc["A1"] = "Detalle por campo (opcional)"
    hc["A1"].font = Font(name=F, bold=True, size=14)
    hc["A2"] = ("Solo para filas con «Campos a modificar» = «Detalle en hoja Campos». Una fila por campo: "
                "«Valor nuevo» (se escribe tal cual) o «Delta» (se suma). Campo como aparece en SIRT, o ruta con « › »: "
                "bloque › campo · tabla · tabla › fila · tabla › fila › campo (ver «Qué hace cada opción»).")
    hc["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    for j, (t, tipo, ancho) in enumerate([("Nombre de la tarifa", "meta", 22), ("Campo (o ruta con «›»)", "obligatorio", 55),
                                          ("Valor nuevo", "opcional", 12), ("Delta", "opcional", 10)], start=1):
        cabeza(hc.cell(5, j, t), tipo)
        hc.column_dimensions[get_column_letter(j)].width = ancho
    for i, fila in enumerate(DETALLE_CAMPOS, start=6):
        for j, v in enumerate(fila, start=1):
            c = hc.cell(i, j, v)
            c.font, c.border = Font(name=F, size=9), BORDE
    hc.freeze_panes = "B6"

    for i, (texto, negrita, tam) in enumerate(DESCRIPCION, start=1):
        guia.cell(i, 1, texto).font = Font(name=F, bold=negrita, size=tam,
                                          color="1F4E78" if negrita and tam == 12 else "000000")
    guia.column_dimensions["A"].width = 115

    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Generado: {SALIDA} ({len(MUESTRAS)} filas de muestra)")


if __name__ == "__main__":
    generar()

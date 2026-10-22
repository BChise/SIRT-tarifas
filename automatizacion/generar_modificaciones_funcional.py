"""
Genera automatizacion/modificaciones_funcional.xlsx: PRUEBA FUNCIONAL pequeña de rectificaciones y
cambios. El objetivo no es volumen sino comprobar que CADA opción del Excel (y de
modificaciones_excel.py / acciones.py) funciona al menos una vez.

- 24 filas que van a SIRT (o se omiten): cada una prueba una combinación distinta (columna
  «Qué se prueba», al final). Usa tarifas de RESERVA (sin base, sin acción, sin modificación
  planificada en modificaciones.xlsx) más empaquetados y licitaciones libres.
- 10 filas con errores de datos A PROPÓSITO (nombres TE-PRUEBA-ERR-…): nunca llegan a SIRT; deben
  quedar en «Error de datos» con el motivo correcto.
- Hoja «Campos» con todas las formas de ruta (campo, bloque › campo, tabla, tabla › fila,
  tabla › fila › campo, lista de la fila con Valor nuevo).

No toca modificaciones.xlsx (la carga grande de 55 filas).

Uso (con el Excel CERRADO):
    python automatizacion\\generar_modificaciones_funcional.py
    python modificaciones_excel.py --excel automatizacion\\modificaciones_funcional.xlsx --solo-listar
    python modificaciones_excel.py --excel automatizacion\\modificaciones_funcional.xlsx --hasta-resumen --limite 1
    python modificaciones_excel.py --excel automatizacion\\modificaciones_funcional.xlsx --hasta-resumen
    python modificaciones_excel.py --excel automatizacion\\modificaciones_funcional.xlsx --si
"""

import sys
from datetime import date, timedelta
from pathlib import Path

AUTO = Path(__file__).resolve().parent
sys.path.insert(0, str(AUTO))
sys.path.insert(0, str(AUTO.parent))
import generar_plantilla_modificaciones as plantilla  # noqa: E402

SALIDA = AUTO / "modificaciones_funcional.xlsx"
EN_5_DIAS = date.today() + timedelta(days=5)
HOY = date.today()
TODO, LISTA, DETALLE, DEFECTO = "Todos los campos", "Solo la lista", "Detalle en hoja Campos", "Por defecto"

# Cada fila: dict con las columnas que se llenan + «prueba» (qué se comprueba).
FILAS = [
    # --- RECTIFICACIÓN: monoproductos ---------------------------------------------------------
    dict(ejecutar="", nombre="TE-TF-CM-010", accion="Rectificación", campos="",
         prueba="Ejecutar vacío = Sí · Campos vacío = Por defecto · Delta vacío = +1 (aviso)"),
    dict(nombre="TE-TFP-CM-009", accion="Rectificación", delta="2", campos=TODO, sello="Rectificación prueba funcional",
         prueba="Delta entero 2 · Todos los campos · texto del sello propio"),
    dict(nombre="TE-IF-CM-011", accion="Rectificación", delta="0.5", campos=LISTA,
         lista="Valor del servicio; Otros datos generales",
         prueba="Delta decimal 0.5 · Solo la lista (2 campos)"),
    dict(nombre="TE-SMC-CM-011", accion="Rectificación", campos=DETALLE,
         prueba="Detalle en rectificación: Valor nuevo + Delta por campo + tabla completa"),
    dict(nombre="TE-SMP-CM-011", accion="Rectificación", campos=TODO, agregar="Sí",
         prueba="Agregar una fila a cada tabla (rectificación)"),
    dict(nombre="TE-TVP-CM-011", accion="Rectificación", campos=TODO, canales="No",
         prueba="TV de paga: Cambiar canales = No"),
    dict(nombre="TE-TVPR-CM-011", accion="Rectificación", campos=DEFECTO,
         prueba="TV de paga: Cambiar canales vacío = Sí (sube el otro Excel de canales)"),
    dict(nombre="TE-OS-CM-010", accion="Rectificación", delta="1", sentido="Beneficioso", campos=DEFECTO,
         lista="Valor del servicio", fecha=EN_5_DIAS,
         prueba="Datos que la rectificación ignora: sentido, fecha y lista → 3 avisos, se aplica igual"),
    # --- CAMBIO: monoproductos ----------------------------------------------------------------
    dict(nombre="TE-TF-CM-011", accion="Cambio", campos=DEFECTO,
         prueba="Delta y sentido vacíos → −1 a todo (escenario mixto, aviso)"),
    dict(nombre="TE-AC-CM-010", accion="Cambio", sentido="Beneficioso", campos=TODO,
         prueba="Beneficioso con Delta vacío (usa 1): precios bajan, beneficios suben"),
    dict(nombre="TE-SMC-CM-010", accion="Cambio", delta="3", sentido="Perjudicial", campos=TODO, ilimitado="Sí",
         prueba="Perjudicial delta 3 · ilimitado → limitado"),
    dict(nombre="TE-SMO-CM-011", accion="Cambio", delta="1", sentido="Beneficioso", campos=TODO, ilimitado="Sí",
         prueba="Beneficioso + ilimitado → limitado (aviso de contradicción, se aplica igual)"),
    dict(nombre="TE-IF-CM-012", accion="Cambio", delta="1", campos=DEFECTO,
         prueba="Delta +1 sin sentido: todo sube (precios y beneficios)"),
    dict(nombre="TE-PIF-CM-012", accion="Cambio", delta="-2", sentido="Perjudicial", campos=DEFECTO, fecha=EN_5_DIAS,
         prueba="Delta negativo con sentido (se usa |−2|) · fecha de ejecución escrita (hoy + 5)"),
    dict(nombre="TE-SMQ-CM-011", accion="Cambio", campos=DETALLE,
         prueba="Detalle en cambio: todas las formas de ruta (ver hoja Campos)"),
    dict(nombre="TE-TVP-CM-012", accion="Cambio", delta="1", sentido="Beneficioso", campos=LISTA,
         lista="Valor del servicio", canales="No", sello="Cambio prueba funcional",
         prueba="Cambio Solo la lista · canales No · sello propio"),
    dict(nombre="TE-SMP-CM-012", accion="Cambio", delta="1", sentido="Beneficioso", campos=TODO, agregar="Sí",
         prueba="Cambio beneficioso + agregar una fila a cada tabla"),
    dict(nombre="TE-PTV-CM-011", accion="Cambio", sentido="Perjudicial", campos=TODO, canales="Sí",
         prueba="Cambio perjudicial TV · Cambiar canales = Sí (escrito)"),
    # --- EMPAQUETADOS ---------------------------------------------------------------------------
    dict(nombre="TE-EMP-CM-113", accion="Rectificación", campos=DEFECTO,
         prueba="Empaquetado: rectificación por defecto (cada pestaña de servicio)"),
    dict(nombre="TE-EMP-CM-117", accion="Cambio", sentido="Beneficioso", campos=TODO,
         prueba="Empaquetado: cambio beneficioso, todos los campos"),
    dict(nombre="TE-EMP-CM-126", accion="Cambio", delta="2", sentido="Perjudicial", campos=DEFECTO,
         prueba="Empaquetado: cambio perjudicial (regla: el último servicio no sube)"),
    # --- LICITACIONES (solo rectificación) ------------------------------------------------------
    dict(nombre="TL-LYC-CM-108", accion="Rectificación", campos=DEFECTO,
         prueba="Licitación 3 servicios: por defecto (vigencia +1 día, ámbito, moneda, valor de contrato)"),
    dict(nombre="TL-LYC-CM-123", accion="Rectificación", campos=TODO,
         prueba="Licitación 3 servicios: todos los campos"),
    # --- OMITIDA ------------------------------------------------------------------------------
    dict(ejecutar="No", nombre="TE-OS-CM-011", accion="Rectificación", campos=DEFECTO,
         prueba="Ejecutar = No → «Omitida» (no toca SIRT)"),
    # --- ERRORES DE DATOS A PROPÓSITO (nunca llegan a SIRT) ---------------------------------------
    dict(ejecutar="Quizás", nombre="TE-PRUEBA-ERR-01", accion="Rectificación", campos=DEFECTO,
         prueba="ERROR esperado: «Ejecutar» distinto de Sí / No"),
    dict(nombre="TE-PRUEBA-ERR-02", accion="Renombre", campos=DEFECTO,
         prueba="ERROR esperado: acción no válida"),
    dict(nombre="TP-PRUEBA-ERR-03", accion="Rectificación", campos=DEFECTO,
         prueba="ERROR esperado: promoción (aún no grabada)"),
    dict(nombre="TL-PRUEBA-ERR-04", accion="Cambio", campos=DEFECTO,
         prueba="ERROR esperado: licitación con cambio (solo rectificación)"),
    dict(nombre="TE-PRUEBA-ERR-05", accion="Rectificación", campos=LISTA,
         prueba="ERROR esperado: «Solo la lista» sin lista"),
    dict(nombre="TE-PRUEBA-ERR-06", accion="Cambio", campos=DETALLE,
         prueba="ERROR esperado: «Detalle en hoja Campos» sin filas en la hoja Campos"),
    dict(nombre="TE-PRUEBA-ERR-07", accion="Cambio", sentido="Neutro", campos=DEFECTO,
         prueba="ERROR esperado: sentido distinto de Beneficioso / Perjudicial"),
    dict(nombre="TE-PRUEBA-ERR-08", accion="Cambio", campos=DEFECTO, fecha=HOY,
         prueba="ERROR esperado: fecha de ejecución que no es posterior a hoy"),
    dict(nombre="TE-PRUEBA-ERR-09", accion="Rectificación", delta="abc", campos=DEFECTO,
         prueba="ERROR esperado: delta que no es número"),
    dict(nombre="TE-PRUEBA-ERR-10", accion="Rectificación", campos="Algunos",
         prueba="ERROR esperado: «Campos a modificar» no válido"),
]

CAMPOS = [  # (tarifa, campo o ruta, valor nuevo, delta, qué se prueba)
    ("TE-SMC-CM-011", "Valor del servicio", "45", "", "campo · Valor nuevo"),
    ("TE-SMC-CM-011", "Datos tope descarga", "", "2", "campo · Delta"),
    ("TE-SMC-CM-011", "Tecnologías y velocidades", "", "1", "tabla completa · Delta"),
    ("TE-SMQ-CM-011", "Valor del servicio", "", "-2", "campo · Delta negativo"),
    ("TE-SMQ-CM-011", "Datos para tethering", "3", "", "campo · Valor nuevo"),
    ("TE-SMQ-CM-011", "Minutos móvil on net › Cantidad de minutos", "300", "", "bloque › campo"),
    ("TE-SMQ-CM-011", "Tecnologías y velocidades › 4G › Descarga máx. contratada", "", "5", "tabla › fila › campo"),
    ("TE-SMQ-CM-011", "Tecnologías y velocidades › 5G", "", "1", "tabla › fila (todos sus números)"),
    ("TE-SMQ-CM-011", "Tecnologías y velocidades › 3G › Tecnología", "4.5G", "", "lista de una fila · Valor nuevo"),
]

COLUMNA = {"ejecutar": "Ejecutar", "nombre": "Nombre de la tarifa", "accion": "Acción",
           "delta": "Delta (vacío: +1 rectificación, −1 cambio)", "sentido": "Cambio: sentido (Beneficioso / Perjudicial)",
           "campos": "Campos a modificar", "lista": "Lista de campos (separados por ;)",
           "agregar": "Agregar una fila a cada tabla", "ilimitado": "Pasar «ilimitado» a limitado",
           "canales": "Cambiar canales (TV de paga)", "fecha": "Cambio: fecha de ejecución",
           "sello": "Texto del sello (vacío: «Rectificación» / «Cambio»)"}


def main() -> None:
    import excel_comun as xc
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if SALIDA.exists() and xc.excel_esta_abierto(SALIDA):
        sys.exit(f"Cierra el Excel antes de generar: {SALIDA}")
    plantilla.SALIDA, plantilla.MUESTRAS, plantilla.DETALLE_CAMPOS = SALIDA, [], []
    plantilla.generar()

    libro = load_workbook(SALIDA)
    hoja = libro["Modificaciones"]
    hoja["A1"] = "Rectificación y cambio de tarifa — PRUEBA FUNCIONAL (todas las opciones)"
    col = {c.value: c.column for c in hoja[5] if c.value}
    lado = Side(style="thin", color="BFBFBF")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    fuente = Font(name="Arial", size=9)
    c_prueba = max(col.values()) + 1
    cab = hoja.cell(5, c_prueba, "Qué se prueba")
    cab.font = Font(name="Arial", bold=True, size=9, color="FFFFFF")
    cab.fill = PatternFill("solid", fgColor="7030A0")
    cab.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    hoja.column_dimensions[get_column_letter(c_prueba)].width = 70
    error = PatternFill("solid", fgColor="FCE4D6")
    for k, f in enumerate(FILAS):
        r = 6 + k
        valores = {"N°": k + 1, "Ejecutar": f.get("ejecutar", "Sí")}
        valores.update({COLUMNA[c]: v for c, v in f.items() if c in COLUMNA and c != "ejecutar"})
        for titulo, valor in valores.items():
            celda = hoja.cell(r, col[titulo], valor if valor != "" else None)
            celda.font, celda.border = fuente, borde
            if isinstance(valor, date):
                celda.number_format = "DD/MM/YYYY"
        celda = hoja.cell(r, c_prueba, f["prueba"])
        celda.font, celda.border = fuente, borde
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        if f["prueba"].startswith("ERROR"):
            celda.fill = error
    hoja.auto_filter.ref = f"A5:{get_column_letter(c_prueba)}{5 + len(FILAS)}"

    hc = libro["Campos"]
    cab = hc.cell(5, 5, "Qué se prueba")
    cab.font = Font(name="Arial", bold=True, size=9, color="FFFFFF")
    cab.fill = PatternFill("solid", fgColor="7030A0")
    hc.column_dimensions["E"].width = 34
    for i, fila in enumerate(CAMPOS, start=6):
        for j, v in enumerate(fila, start=1):
            c = hc.cell(i, j, v or None)
            c.font, c.border = fuente, borde
    libro.save(SALIDA)
    sigue = sum(1 for f in FILAS if not f["prueba"].startswith("ERROR") and f.get("ejecutar") != "No")
    print(f"Listo: {SALIDA.name} · {len(FILAS)} filas: {sigue} van a SIRT, 1 omitida, "
          f"{sum(1 for f in FILAS if f['prueba'].startswith('ERROR'))} errores de datos a propósito · "
          f"{len(CAMPOS)} filas en la hoja Campos")


if __name__ == "__main__":
    main()

"""
Genera automatizacion/propositos_tarifas.xlsx: una fila por tarifa de prueba (monoproductos,
empaquetados, licitaciones y promociones) con su PROPÓSITO, armado a partir de los Excel de
automatizacion/ (registro, empaquetados, licitaciones, promociones, acciones y modificaciones) y del
estado en resultados/ (registrada / registro definitivo).

Uso (con los Excel cerrados no es necesario: solo los lee):
    python automatizacion\\generar_propositos.py
"""

import csv
import sys
from collections import defaultdict
from pathlib import Path

PROYECTO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROYECTO))
import excel_comun as xc  # noqa: E402

AUTO = PROYECTO / "automatizacion"
SALIDA = AUTO / "propositos_tarifas.xlsx"
SERVICIOS = ["TF", "TFP", "OTF", "PTF", "IF", "PIF", "TVP", "TVPR", "PTV", "SMP", "SMC", "SMO", "SMQ", "OS", "AC"]
NOMBRES = {"TF": "Telefonía fija postpago", "TFP": "Telefonía fija prepago", "OTF": "Otras tarifas de telefonía fija",
           "PTF": "Paquetes adicionales de telefonía fija", "IF": "Internet fijo postpago",
           "PIF": "Paquetes adicionales de internet fijo", "TVP": "TV de paga postpago", "TVPR": "TV de paga prepago",
           "PTV": "Paquetes adicionales de TV de paga", "SMP": "Móvil postpago", "SMC": "Móvil control",
           "SMO": "Móvil otros", "SMQ": "Móvil paquete", "OS": "Otros servicios", "AC": "Arrendamiento de circuitos"}


def libro(nombre):
    ruta = AUTO / nombre
    return xc.abrir_libro(ruta, True) if ruta.exists() else None


def filas(nombre_libro, hoja, columna):
    lb = libro(nombre_libro)
    if lb is None or hoja not in lb.sheetnames:
        return []
    return [f for _, f in xc.Hoja(lb[hoja]).filas(columna)]


def estados() -> dict:
    """nombre → 'Registro definitivo' / 'Registrada' (según resultados/)."""
    salida = {}
    for ruta in (PROYECTO / "resultados").glob("*/resultados_*.csv"):
        with ruta.open("r", newline="", encoding="utf-8-sig") as e:
            for f in csv.DictReader(e, delimiter=";"):
                if f.get("resultado") == "EXITOSO":
                    salida.setdefault(f["nombre"], "Registrada")
    for ruta in (PROYECTO / "resultados").glob("registro_definitivo_*.csv"):
        with ruta.open("r", newline="", encoding="utf-8-sig") as e:
            for f in csv.DictReader(e, delimiter=";"):
                if f.get("resultado") == "EXITOSO":
                    salida[f["nombre"]] = "Registro definitivo"
    return salida


def main() -> None:
    estado = estados()
    base_de = defaultdict(list)                      # tarifa → paquetes / promociones que la usan
    emp_info, prm_info, lyc_info = {}, {}, {}
    for f in filas("empaquetados.xlsx", "Empaquetados", "Nombre del paquete"):
        n = xc.texto(f["Nombre del paquete"])
        tarifas = [xc.texto(f.get(f"Tarifa {k}")) for k in range(1, 7) if xc.texto(f.get(f"Tarifa {k}"))]
        emp_info[n] = (tarifas, xc.texto(f.get("Otros datos generales")), xc.ejecutar_fila(f))
        for t in tarifas:
            base_de[t].append(n)
    for f in filas("promociones.xlsx", "Promociones", "Nombre de la promoción"):
        n = xc.texto(f["Nombre de la promoción"])
        tarifas = [xc.texto(v) for t, v in f.items() if t.startswith("Tarifa base ") and xc.texto(v)]
        prm_info[n] = (tarifas, xc.texto(f.get("Tipo de variación preferido")) or "orden normal",
                       xc.texto(f.get("Atributos a promocionar")))
        for t in tarifas:
            base_de[t].append(n)
    for f in filas("registro_licitaciones.xlsx", "Licitaciones", "Nombre de la tarifa"):
        servicios = [s for s in ("Arrendamiento de circuitos", "Otros servicios", "Televisión de paga", "Internet Fijo",
                                 "Servicio Móvil", "Telefonía Fija") if xc.es_si(f.get(f"{s}: incluir"))]
        lyc_info[xc.texto(f["Nombre de la tarifa"])] = servicios
    acciones_de = defaultdict(list)                  # una tarifa puede tener varias (promociones mixtas)
    for f in filas("acciones_basicas.xlsx", "Acciones", "Nombre de la tarifa"):
        acciones_de[xc.texto(f["Nombre de la tarifa"])].append(xc.texto(f["Acción"]))
    accion = {n: ", ".join(a) for n, a in acciones_de.items()}
    modif = {}
    for f in filas("modificaciones.xlsx", "Modificaciones", "Nombre de la tarifa"):
        sentido = xc.texto(f.get("Cambio: sentido (Beneficioso / Perjudicial)"))
        texto_m = xc.texto(f["Acción"]) + (f" {sentido.lower()}" if sentido else "")
        modif[xc.texto(f["Nombre de la tarifa"])] = f"{texto_m} ({xc.texto(f.get('Campos a modificar')).lower()})"

    def resumen(nombre, base):
        partes = []
        if base:
            partes.append(base)
        if nombre in base_de:
            emp = [x for x in base_de[nombre] if x.startswith("TE-EMP")]
            prm = [x for x in base_de[nombre] if x.startswith("TP-")]
            if emp:
                partes.append(f"base de {len(emp)} empaquetado(s)")
            if prm:
                partes.append(f"base de {len(prm)} promoción(es)")
        if nombre in accion:
            partes.append(accion[nombre].lower())
        if nombre in modif:
            partes.append(modif[nombre].split(" (")[0].lower())
        return "; ".join(partes) or "reserva (sin uso asignado)"

    salida = []
    reg = libro("registro_tarifas_establecidas.xlsx")
    for codigo in SERVICIOS:
        if reg is None or codigo not in reg.sheetnames:
            continue
        for k, f in enumerate(xc.Hoja(reg[codigo]).filas("Nombre de la tarifa")):
            f = f[1]
            n = xc.texto(f["Nombre de la tarifa"])
            base = "testigo: queda tal como se registró (sin acciones)" if k == 0 else ""
            salida.append(("Monoproducto", codigo, NOMBRES[codigo], n, xc.texto(f.get("Alcance de la tarifa")),
                           resumen(n, base), ", ".join(base_de.get(n, [])), accion.get(n, ""), modif.get(n, ""),
                           estado.get(n, "Pendiente")))
    for n, (tarifas, desc, ej) in emp_info.items():
        codigos = [t.split("-")[1] for t in tarifas]
        base = desc or f"{len(tarifas)} servicios: " + " + ".join(codigos)
        if ej == "no":
            base += " (Ejecutar = No)"
        salida.append(("Empaquetado", "EMP", " + ".join(codigos), n, "", resumen(n, base),
                       ", ".join(base_de.get(n, [])), accion.get(n, ""), modif.get(n, ""), estado.get(n, "Pendiente")))
    for n, servicios in lyc_info.items():
        base = f"{len(servicios)} servicio(s): " + " + ".join(servicios)
        salida.append(("Licitación", "LYC", " + ".join(servicios), n, "", resumen(n, base), "",
                       accion.get(n, ""), modif.get(n, ""), estado.get(n, "Pendiente")))
    for n, (tarifas, preferido, modo) in prm_info.items():
        if len(tarifas) == 1 and tarifas[0].startswith("TE-EMP"):
            base = f"promoción de un empaquetado ({tarifas[0]})"
        elif len(tarifas) == 1:
            base = f"promoción de un servicio, atributos «{modo}» con preferido {preferido}"
        elif any(t.startswith("TE-EMP") for t in tarifas):
            base = f"mixta: {sum(not t.startswith('TE-EMP') for t in tarifas)} monoproductos + " \
                   f"{sum(t.startswith('TE-EMP') for t in tarifas)} empaquetados (SIRT: sin atributos en común)"
        else:
            base = f"{len(tarifas)} monoproductos de servicios distintos (SIRT: sin atributos en común)"
        salida.append(("Promoción", "PRM", " + ".join(tarifas)[:200], n, "", resumen(n, base), "",
                       accion.get(n, ""), modif.get(n, ""), estado.get(n, "Pendiente")))

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    h = wb.active
    h.title = "Propósitos"
    h["A1"] = "Propósito de cada tarifa de prueba"
    h["A1"].font = Font(name="Arial", bold=True, size=14)
    h["A2"] = ("Generado desde los Excel de automatizacion/ con: python automatizacion\\generar_propositos.py. "
               "Vuelve a generarlo si cambian los Excel.")
    h["A2"].font = Font(name="Arial", italic=True, size=9, color="595959")
    cab = [("Tipo", 13), ("Código", 8), ("Servicio(s)", 34), ("Nombre de la tarifa", 22), ("Alcance", 20),
           ("Propósito", 70), ("Es base de", 40), ("Acción básica", 14), ("Modificación", 34), ("Estado", 18)]
    for j, (t, ancho) in enumerate(cab, start=1):
        c = h.cell(4, j, t)
        c.font, c.fill = Font(name="Arial", bold=True, size=9, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.column_dimensions[get_column_letter(j)].width = ancho
    colores = {"Monoproducto": "FFFFFF", "Empaquetado": "F2F2F2", "Licitación": "FFFFFF", "Promoción": "F2F2F2"}
    for i, fila in enumerate(salida, start=5):
        for j, v in enumerate(fila, start=1):
            c = h.cell(i, j, v)
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="top", wrap_text=j in (3, 6, 7, 9))
            c.fill = PatternFill("solid", fgColor=colores[fila[0]])
    h.freeze_panes = "E5"
    h.auto_filter.ref = f"A4:{get_column_letter(len(cab))}{4 + len(salida)}"
    wb.save(SALIDA)
    from collections import Counter
    print(f"Listo: {SALIDA.name} con {len(salida)} tarifas (" +
          ", ".join(f"{v} {k.lower()}" for k, v in Counter(f[0] for f in salida).items()) + ")")


if __name__ == "__main__":
    main()

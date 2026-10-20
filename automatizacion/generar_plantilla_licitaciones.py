"""
Genera automatizacion/registro_licitaciones.xlsx (BORRADOR, sin conectar con los
scripts): el Excel para registrar tarifas de LICITACIÓN Y CONTRATOS.

Una fila = una tarifa. Los servicios se eligen con una columna Sí/No por servicio (sin
escribir nombres) y cada servicio elegido lleva su «Detalle del servicio» (paso 2).

Uso (desde la carpeta del proyecto):
    python automatizacion/generar_plantilla_licitaciones.py
"""

from pathlib import Path

PROYECTO = Path(__file__).resolve().parent.parent
SALIDA = PROYECTO / "automatizacion" / "registro_licitaciones.xlsx"
FILAS_CON_LISTAS = 1000

# Nombres exactos en SIRT (servicios.json de licitacion_contratos).
SERVICIOS = ["Arrendamiento de circuitos", "Otros servicios", "Televisión de paga",
             "Internet Fijo", "Servicio Móvil", "Telefonía Fija"]
LISTAS = {"Sí/No": ["Sí", "No"], "Ámbito": ["Público", "Privado"], "Moneda": ["Soles", "Dolares", "Euro"]}

# (título, tipo, lista, ancho, sección)   tipo: meta | obligatorio | opcional | script
COLUMNAS = [
    ("N°", "meta", None, 5, ""), ("Ejecutar", "meta", "Sí/No", 9, ""),
    ("Nombre de la tarifa", "meta", None, 22, ""), ("Registro definitivo", "meta", "Sí/No", 11, ""),
    ("Inicio de vigencia", "obligatorio", None, 12, "Datos generales"),
    ("Duración del contrato (días)", "obligatorio", None, 12, "Datos generales"),
    ("Identificador de contrato", "obligatorio", None, 16, "Datos generales"),
    ("Ámbito", "obligatorio", "Ámbito", 10, "Datos generales"),
    ("Valor de contrato", "obligatorio", None, 12, "Características específicas"),
    ("Moneda", "obligatorio", "Moneda", 10, "Características específicas"),
    ("Documento de Empresa", "obligatorio", None, 15, "Características adicionales"),
    ("Nombre de empresa", "obligatorio", None, 24, "Características adicionales"),
    ("Descripción del Servicio", "obligatorio", None, 30, "Características adicionales"),
    ("Atributos de la tarifa", "obligatorio", None, 30, "Características adicionales"),
    ("Ámbito geográfico", "obligatorio", None, 20, "Características adicionales"),
    ("Tarifa instalación (solo con Otros servicios)", "opcional", None, 14, "Características adicionales"),
] + [c for s in SERVICIOS for c in (
    (f"{s}: incluir", "obligatorio", "Sí/No", 11, "Servicios (paso 2: un detalle por servicio incluido)"),
    (f"{s}: detalle del servicio", "opcional", None, 26, "Servicios (paso 2: un detalle por servicio incluido)"),
)] + [
    ("Cantidad de servicios", "calculo", None, 10, ""),
    ("Estado", "script", None, 13, "Resultado"), ("Código SIRT", "script", None, 18, "Resultado"),
    ("Fecha de registro", "script", None, 17, "Resultado"), ("Detalle", "script", None, 50, "Resultado"),
]

# Muestras: 1, 3 y 6 servicios (como los grupos que se definieron).
MUESTRAS = [
    {"servicios": ["Internet Fijo"], "ambito": "Público", "moneda": "Soles", "valor": "1500.50"},
    {"servicios": ["Televisión de paga", "Servicio Móvil", "Telefonía Fija"], "ambito": "Privado",
     "moneda": "Dolares", "valor": "3250.75"},
    {"servicios": SERVICIOS, "ambito": "Público", "moneda": "Euro", "valor": "9800"},
]


def generar() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    lado = Side(style="thin", color="BFBFBF")
    BORDE = Border(left=lado, right=lado, top=lado, bottom=lado)
    RELLENO = {"meta": ("375623", "FFFFFF"), "obligatorio": ("1F4E78", "FFFFFF"), "opcional": ("BDD7EE", "1F1F1F"),
               "script": ("808080", "FFFFFF"), "calculo": ("E2EFDA", "1F1F1F")}
    CELDA_SCRIPT = PatternFill("solid", fgColor="EDEDED")
    SECCION = PatternFill("solid", fgColor="DDEBF7")

    wb = Workbook()
    guia = wb.active
    guia.title = "Instrucciones"
    hoja = wb.create_sheet("Licitaciones")
    listas = wb.create_sheet("Listas")
    listas.sheet_state = "hidden"

    hoja["A1"] = "Licitación y Contratos"
    hoja["A1"].font = Font(name=F, bold=True, size=14)
    hoja["A2"] = ("Una fila por tarifa. Marca con «Sí» los servicios que incluye y escribe el detalle de cada uno. "
                  "Fin de vigencia = inicio + duración − 1.")
    hoja["A2"].font = Font(name=F, italic=True, size=9, color="595959")
    for j, (texto, clave) in enumerate([("Obligatorio", "obligatorio"), ("Opcional", "opcional"),
                                        ("Calculado", "calculo"), ("Lo llena el script", "script")], start=4):
        fondo, color = RELLENO[clave]
        c = hoja.cell(3, j, texto)
        c.fill, c.border = PatternFill("solid", fgColor=fondo), BORDE
        c.font = Font(name=F, bold=True, size=9, color=color)
        c.alignment = Alignment(horizontal="center")

    secciones = [c[4] for c in COLUMNAS]
    for j, s in enumerate(secciones, start=1):
        c = hoja.cell(4, j, s if (s and (j == 1 or s != secciones[j - 2])) else None)
        c.fill, c.border = SECCION, BORDE
        c.font = Font(name=F, bold=True, size=8, color="1F4E78")
    for j, (titulo, tipo, _, ancho, _) in enumerate(COLUMNAS, start=1):
        fondo, color = RELLENO[tipo]
        c = hoja.cell(5, j, titulo)
        c.fill, c.border = PatternFill("solid", fgColor=fondo), BORDE
        c.font = Font(name=F, bold=True, size=9, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        hoja.column_dimensions[get_column_letter(j)].width = ancho
    hoja.row_dimensions[5].height = 54

    titulos = [c[0] for c in COLUMNAS]
    col = {t: k + 1 for k, t in enumerate(titulos)}
    for i, m in enumerate(MUESTRAS, start=6):
        n = i - 5
        codigos = "+".join({"Arrendamiento de circuitos": "AC", "Otros servicios": "OS", "Televisión de paga": "TVP",
                            "Internet Fijo": "IF", "Servicio Móvil": "SM", "Telefonía Fija": "TF"}[s]
                           for s in m["servicios"])
        valores = {
            "N°": n, "Ejecutar": "Sí", "Nombre de la tarifa": f"TL-LYC-CM-{n:03d}", "Registro definitivo": "Sí",
            "Inicio de vigencia": "", "Duración del contrato (días)": 365, "Identificador de contrato": f"CTR-CM-{n:03d}",
            "Ámbito": m["ambito"], "Valor de contrato": m["valor"], "Moneda": m["moneda"],
            "Documento de Empresa": f"{20100000000 + n}", "Nombre de empresa": f"Empresa contratante de prueba {n:03d} SAC",
            "Descripción del Servicio": f"Servicio licitado de prueba ({codigos})",
            "Atributos de la tarifa": f"Atributos de la tarifa de prueba {n:03d}",
            "Ámbito geográfico": "Lima Metropolitana",
            "Tarifa instalación (solo con Otros servicios)": "10" if "Otros servicios" in m["servicios"] else "",
        }
        for s in SERVICIOS:
            incluido = s in m["servicios"]
            valores[f"{s}: incluir"] = "Sí" if incluido else "No"
            valores[f"{s}: detalle del servicio"] = f"Detalle de {s.lower()} para la tarifa {n:03d}" if incluido else ""
        for t, v in valores.items():
            c = hoja.cell(i, col[t], v)
            c.font, c.border = Font(name=F, size=9), BORDE
            c.alignment = Alignment(vertical="top", wrap_text=col[t] > 3,
                                    horizontal="center" if t in ("N°", "Ejecutar", "Registro definitivo")
                                    or t.endswith(": incluir") else None)
        for t in ("Estado", "Código SIRT", "Fecha de registro", "Detalle"):
            c = hoja.cell(i, col[t])
            c.fill, c.border = CELDA_SCRIPT, BORDE
    # «Cantidad de servicios»: cuenta los «Sí» de las columnas «incluir».
    incluir = [get_column_letter(col[f"{s}: incluir"]) for s in SERVICIOS]
    c_cant = col["Cantidad de servicios"]
    for i in range(6, 6 + FILAS_CON_LISTAS):
        formula = "+".join(f'({l}{i}="Sí")' for l in incluir)
        c = hoja.cell(i, c_cant, f'=IF(C{i}="","",{formula})')
        c.fill = PatternFill("solid", fgColor="E2EFDA")
        c.border, c.font = BORDE, Font(name=F, size=9)
        c.alignment = Alignment(horizontal="center", vertical="top")
    for i in range(6, 6 + FILAS_CON_LISTAS):
        hoja.cell(i, col["Inicio de vigencia"]).number_format = "DD/MM/YYYY"

    # Listas desplegables desde la hoja oculta «Listas».
    rangos = {}
    for k, (nombre, opciones) in enumerate(LISTAS.items(), start=1):
        letra = get_column_letter(k)
        for f, o in enumerate(opciones, start=1):
            listas[f"{letra}{f}"] = o
        rangos[nombre] = f"=Listas!${letra}$1:${letra}${len(opciones)}"
    for titulo, _, lista, _, _ in COLUMNAS:
        if lista:
            dv = DataValidation(type="list", formula1=rangos[lista], allow_blank=True)
            dv.showErrorMessage, dv.errorTitle, dv.error = True, "Valor no válido", "Elige una opción de la lista."
            hoja.add_data_validation(dv)
            letra = get_column_letter(col[titulo])
            dv.add(f"{letra}6:{letra}{5 + FILAS_CON_LISTAS}")
    numero = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
    numero.showErrorMessage, numero.error = True, "La duración es un número entero de días (mayor que 0)."
    hoja.add_data_validation(numero)
    letra = get_column_letter(col["Duración del contrato (días)"])
    numero.add(f"{letra}6:{letra}{5 + FILAS_CON_LISTAS}")

    hoja.freeze_panes = "D6"
    hoja.auto_filter.ref = f"A5:{get_column_letter(len(COLUMNAS))}{5 + len(MUESTRAS)}"

    textos = [
        ("Registro de Licitación y Contratos (BORRADOR)", True, 14),
        ("Ubicación: automatizacion/registro_licitaciones.xlsx.", False, 9),
        ("", False, 10),
        ("Cómo se usa", True, 11),
        ("1. Una fila por tarifa (agrega las que necesites). «Ejecutar» = No: la fila se deja sin tocar.", False, 10),
        ("2. Servicios: pon «Sí» en «<servicio>: incluir» para cada servicio de la tarifa (1 a 6) y escribe su", False, 10),
        ("   «detalle del servicio» (es lo que pide SIRT en cada pestaña del paso 2). «Cantidad de servicios» se calcula sola.", False, 10),
        ("3. Fechas: «Inicio de vigencia» vacío = día del registro. El fin de vigencia lo calcula el script:", False, 10),
        ("   inicio + duración − 1 (SIRT exige que la duración cuadre con las fechas).", False, 10),
        ("4. Guarda, cierra el Excel y corre el script. Las columnas grises mostrarán el resultado de cada fila.", False, 10),
        ("", False, 10),
        ("Valores", True, 11),
        ("• Ámbito: Público / Privado. Moneda: Soles / Dolares / Euro (como en SIRT).", False, 10),
        ("• Valor de contrato: número, hasta 4 decimales.", False, 10),
        ("• Documento de Empresa: solo dígitos (máx. 15). Textos de características: hasta 4000 caracteres.", False, 10),
        ("• Tarifa instalación: solo si la tarifa incluye «Otros servicios» (único que la registra).", False, 10),
        ("• Registro definitivo: sin disponibilidad comercial (los servicios de licitación no la piden).", False, 10),
        ("", False, 10),
        ("Después del registro", True, 11),
        ("La fe de erratas y la rectificación de estas tarifas se hacen con los Excel de acciones (fe de erratas: acciones_basicas.xlsx).", False, 10),
        ("", False, 10),
        ("Colores", True, 11),
        ("Verde oscuro: datos de la fila. Azul oscuro: obligatorio. Azul claro: opcional. Verde claro: calculado. "
         "Gris: lo llena el script.", False, 10),
    ]
    for i, (t, negrita, tam) in enumerate(textos, start=1):
        guia.cell(i, 1, t).font = Font(name=F, bold=negrita, size=tam)
    guia.column_dimensions["A"].width = 120

    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Generado: {SALIDA} ({len(MUESTRAS)} tarifas de muestra)")


if __name__ == "__main__":
    generar()

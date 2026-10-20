"""
Genera automatizacion/registro_tarifas_establecidas.xlsx (BORRADOR, aún sin conectar
con los scripts): el Excel donde se escriben las tarifas establecidas a registrar.

- Una hoja por servicio. Cada fila = una tarifa (1, 2, 3… las que se escriban).
- Cabecera y columnas «N°» / «Nombre de la tarifa» fijas al desplazarse.
- Cabecera azul oscuro = atributo obligatorio; azul claro = opcional
  (según «esObligatorio» del JSON del paso 2 de SIRT; los del paso 1, por lo observado).
- Columnas grises del final: las llenará el script (estado, código SIRT, detalle).

Uso (desde la carpeta del proyecto):
    python automatizacion/generar_plantilla_registro.py
"""

import importlib
import json
import re
import sys
from pathlib import Path

PROYECTO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROYECTO))
sys.path.insert(0, str(PROYECTO / "datos"))

import comun  # noqa: E402
import generar_excel as base  # noqa: E402   (reutiliza atributos() y sus etiquetas)
from app import catalogo  # noqa: E402

SALIDA = PROYECTO / "automatizacion" / "registro_tarifas_establecidas.xlsx"
FILAS_EJEMPLO = 3
FILAS_CON_LISTAS = 1000                      # filas con listas desplegables preparadas

META = ["N°", "Ejecutar", "Nombre de la tarifa", "Registro definitivo"]
SALIDA_SCRIPT = ["Estado", "Código SIRT", "Fecha de registro", "Detalle"]

# Columnas del paso 1 (no están en el JSON del paso 2): obligatoriedad observada.
PASO_1 = {"Alcance": True, "Tarifa de instalación": True, "Otros datos generales": True}
# Nombre de la columna → nombre del apartado en el JSON de SIRT.
ALIAS = {
    "Tarifas fijo LDI (tabla)": "Tarifa por minuto a fijo LDI",
    "Tarifas móvil LDI (tabla)": "Tarifa por minuto a móvil LDI",
    "Velocidad navegacion": "Velocidades de navegación",
    "Tecnologías y velocidades (tabla)": "Tecnologías y velocidades",
}
# Columnas sin apartado en el JSON (listas del paso 2 de IF, que no tiene JSON propio, etc.).
SIN_JSON = {"Periodicidad": True, "Tipo de acceso": True, "Tecnología (IF)": True,
            "Internet Ilimitado": True, "Agregar degradacion": False, "Velocidad degradacion": False}
# Nombres más claros para algunas columnas internas.
RENOMBRAR = {"Velocidad navegacion": "Velocidades de navegación (tabla)",
             "Agregar degradacion": "¿Agregar velocidades con degradación?",
             "Velocidad degradacion": "Velocidades con degradación (tabla)"}
ALCANCES = ["Residencial", "Comercial", "Residencial + Comercial"]


def obligatoriedad(archivo: str) -> dict:
    ruta = PROYECTO / "grabaciones" / archivo / "paso2.json"
    if not ruta.exists() or ruta.stat().st_size == 0:
        return {}
    datos = json.loads(ruta.read_text(encoding="utf-8"))["data"]
    salida = {}
    for seccion in datos.get("secciones") or []:
        for apartado in seccion.get("apartados") or []:
            salida[catalogo._sin_tildes(apartado["nombre"])] = bool(apartado.get("esObligatorio"))
            for sub in apartado.get("subApartados") or []:
                salida[catalogo._sin_tildes(sub["nombre"])] = bool(sub.get("esObligatorio"))
    return salida


def es_obligatorio(columna: str, tabla: dict) -> bool:
    if columna in PASO_1:
        return PASO_1[columna]
    nombre = ALIAS.get(columna, re.sub(r"\s*\(.*?\)\s*$", "", columna))
    clave = catalogo._sin_tildes(nombre)
    if clave in tabla:
        return tabla[clave]
    return SIN_JSON.get(columna, True)


def generar() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = "Arial"
    lado = Side(style="thin", color="BFBFBF")
    BORDE = Border(left=lado, right=lado, top=lado, bottom=lado)
    OBLIG = PatternFill("solid", fgColor="1F4E78")      # cabecera: obligatorio
    OPCIONAL = PatternFill("solid", fgColor="BDD7EE")   # cabecera: opcional
    META_F = PatternFill("solid", fgColor="375623")     # cabecera: datos de la fila
    SCRIPT = PatternFill("solid", fgColor="808080")     # cabecera: lo llena el script
    CELDA_SCRIPT = PatternFill("solid", fgColor="EDEDED")

    wb = Workbook()
    guia = wb.active
    guia.title = "Instrucciones"
    resumen = []

    for servicio in catalogo.servicios():
        modulo = importlib.import_module(f"servicios.{servicio['archivo']}")
        casos = comun.generar_casos(modulo.DIMENSIONES, modulo.CODIGO, "REG")
        origen = catalogo.RESPALDO.get(servicio["codigo"], (servicio["archivo"],))[0] \
            if not (PROYECTO / "grabaciones" / servicio["archivo"] / "paso2.json").exists() else servicio["archivo"]
        tabla_oblig = obligatoriedad(origen)

        filas = []
        for k in range(1, FILAS_EJEMPLO + 1):
            buscado = {a.strip().lower() for a in ALCANCES[(k - 1) % 3].split("+")}
            caso = next((c for c in casos if {a.strip().lower() for a in
                         comun.texto_valor(c["valores"].get("alcance", "")).split("+")} == buscado), casos[0])
            filas.append((k, f"TE-{modulo.CODIGO}-REG-{k:03d}", base.atributos(modulo, caso)))
        columnas = list(filas[0][2])
        oblig = {c: es_obligatorio(c, tabla_oblig) for c in columnas}

        hoja = wb.create_sheet(modulo.CODIGO)
        hoja["A1"] = f"{modulo.CODIGO} · {servicio['nombre']}"
        hoja["A1"].font = Font(name=F, bold=True, size=14)
        hoja["A2"] = ("Una fila por tarifa a registrar (agrega las que necesites debajo). "
                      "Las columnas grises del final las llena el script.")
        hoja["A2"].font = Font(name=F, italic=True, size=9, color="595959")
        for j, (texto, relleno, color) in enumerate(
                [("Obligatorio", OBLIG, "FFFFFF"), ("Opcional", OPCIONAL, "1F1F1F"),
                 ("Lo llena el script", SCRIPT, "FFFFFF")], start=3):
            c = hoja.cell(3, j, texto)
            c.fill, c.font, c.border = relleno, Font(name=F, bold=True, size=9, color=color), BORDE
            c.alignment = Alignment(horizontal="center")

        cabecera = META + [base_nombre for base_nombre in columnas] + SALIDA_SCRIPT
        for j, titulo in enumerate(cabecera, start=1):
            visible = RENOMBRAR.get(titulo, titulo)
            c = hoja.cell(5, j, visible)
            if j <= len(META):
                relleno, color = META_F, "FFFFFF"
            elif titulo in SALIDA_SCRIPT and j > len(META) + len(columnas):
                relleno, color = SCRIPT, "FFFFFF"
            elif oblig.get(titulo, True):
                relleno, color = OBLIG, "FFFFFF"
            else:
                relleno, color = OPCIONAL, "1F1F1F"
            c.fill, c.border = relleno, BORDE
            c.font = Font(name=F, bold=True, size=9, color=color)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        hoja.row_dimensions[5].height = 45

        for i, (k, nombre, attrs) in enumerate(filas, start=6):
            valores = [k, "Sí", nombre, "Sí"] + [attrs.get(c, "") for c in columnas] + [""] * len(SALIDA_SCRIPT)
            for j, v in enumerate(valores, start=1):
                c = hoja.cell(i, j, v)
                c.font = Font(name=F, size=9)
                c.border = BORDE
                c.alignment = Alignment(vertical="top", wrap_text=j > 3,
                                        horizontal="center" if j in (1, 2, 4) else None)
                if j > len(META) + len(columnas):
                    c.fill = CELDA_SCRIPT

        ultima = 5 + FILAS_CON_LISTAS
        validaciones = [(2, '"Sí,No"'), (4, '"Sí,No"')]
        if "Alcance" in columnas:
            validaciones.append((len(META) + columnas.index("Alcance") + 1, '"' + ",".join(ALCANCES) + '"'))
        for col, opciones in validaciones:
            dv = DataValidation(type="list", formula1=opciones, allow_blank=True)
            hoja.add_data_validation(dv)
            letra = get_column_letter(col)
            dv.add(f"{letra}6:{letra}{ultima}")

        anchos = [5, 9, 24, 11] + [max(14, min(38, len(RENOMBRAR.get(c, c)) + 4)) for c in columnas] + [12, 18, 16, 40]
        for j, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(j)].width = ancho
        hoja.freeze_panes = "D6"                         # cabecera y N° / Ejecutar / Nombre siempre visibles
        hoja.auto_filter.ref = f"A5:{get_column_letter(len(cabecera))}{5 + len(filas)}"
        resumen.append((modulo.CODIGO, servicio["nombre"], len(columnas),
                        sum(1 for v in oblig.values() if v), sum(1 for v in oblig.values() if not v)))

    textos = [
        ("Registro de tarifas establecidas (BORRADOR)", True, 14),
        ("", False, 10),
        ("Cómo se usa", True, 11),
        ("1. Abre la hoja del servicio (código en la pestaña).", False, 10),
        ("2. Escribe una fila por tarifa a registrar: 1, 2, 3… las que quieras, sin dejar filas vacías en medio.", False, 10),
        ("   «Ejecutar» = Sí: la fila se procesa; = No: se deja en el Excel pero no se toca (estado «Omitida»).", False, 10),
        ("3. Completa sus atributos. El nombre de la tarifa debe ser único.", False, 10),
        ("4. Guarda el archivo y corre el script de registro.", False, 10),
        ("5. El script escribirá en las columnas grises el estado de cada fila (Registrada, Definitivo o Error).", False, 10),
        ("", False, 10),
        ("Colores de la cabecera", True, 11),
        ("Verde: datos de la fila (N°, ejecutar, nombre, registro definitivo).", False, 10),
        ("Azul oscuro: atributo obligatorio en SIRT. Azul claro: opcional (puede quedar vacío).", False, 10),
        ("Gris: lo llena el script; no escribir ahí.", False, 10),
        ("", False, 10),
        ("Formato de algunos valores (por ahora)", True, 11),
        ("Listas de varias opciones: separadas por « + » (p. ej. Residencial + Comercial).", False, 10),
        ("Minutos/SMS: «Ilimitado» o «Limitado · minutos: 100 · tarifa: 0.10».", False, 10),
        ("Tablas (países, velocidades, bolsas): resumen en una celda; en la versión oficial irán en una hoja aparte.", False, 10),
        ("", False, 10),
        ("Servicios", True, 11),
    ]
    for i, (t, negrita, tam) in enumerate(textos, start=1):
        guia.cell(i, 1, t).font = Font(name=F, bold=negrita, size=tam)
    fila = len(textos) + 1
    for j, titulo in enumerate(["Hoja", "Servicio", "Atributos", "Obligatorios", "Opcionales"], start=1):
        c = guia.cell(fila, j, titulo)
        c.font, c.fill = Font(name=F, bold=True, color="FFFFFF"), OBLIG
    for codigo, nombre, n, o, op in resumen:
        fila += 1
        for j, v in enumerate([codigo, nombre, n, o, op], start=1):
            guia.cell(fila, j, v).font = Font(name=F, size=10)
    guia.column_dimensions["A"].width = 12
    guia.column_dimensions["B"].width = 58
    for col in "CDE":
        guia.column_dimensions[col].width = 12

    SALIDA.parent.mkdir(exist_ok=True)
    wb.save(SALIDA)
    print(f"Generado: {SALIDA}")
    for codigo, nombre, n, o, op in resumen:
        print(f"  {codigo:<5} {n:>3} atributos ({o} obligatorios, {op} opcionales)")


if __name__ == "__main__":
    generar()
